# -*- coding: utf-8 -*-
"""
pile_engine_v5.py — Μηχανή παραγωγής Excel κοστολόγησης πασσάλων (RULES_v5 / AGENTS_v5)
ΘΕΜΕΛΙΟΤΕΧΝΙΚΗ ΛΤΔ

Χρήση:
    python pile_engine_v5.py <project_spec.json> [--out <file.xlsx>] [--no-verify]

Βασικές αρχές (v5.7 — Αρχιτεκτονική 3 Φύλλων, 100% Δυναμικές Σημειώσεις Παντού):
  1. ΚΑΜΙΑ διεύθυνση κελιού δεν γράφεται με το χέρι. Κάθε τιμή καταχωρείται με ΟΝΟΜΑ
     (cell registry) και οι φόρμουλες γράφονται ως πρότυπα "={L_pile}+{dh}".
  2. IRON RULE ΔΥΝΑΜΙΚΩΝ ΣΗΜΕΙΩΣΕΩΝ: Όλες οι σημειώσεις και επεξηγήσεις (Φύλλο 1, Φύλλο 2 Ενότητα 2Β,
     Φύλλο 2 Ενότητα 5 Drawing Audit) είναι 100% ζωντανές φόρμουλες Excel συνδεδεμένες με τα κελιά!
  3. 3 Φύλλα:
     - Φύλλο 1: «Οικονομική Προσφορά» (Καθαρό τιμολόγιο πελάτη + 100% δυναμικές σημειώσεις)
     - Φύλλο 2: «Τεχνικοί Υπολογισμοί» (Δεδομένα, Σμαραγδί Πίνακας, Ενότητα 2Β με ανθρώπινες
       επεξηγήσεις για τον μάστρο, Σταθερές, Αναλυτικά Μαθηματικά, Drawing Audit)
     - Φύλλο 3: «Τιμές & Έλεγχος» (Ο ΜΟΝΑΔΙΚΟΣ ΑΡΧΗΓΟΣ ΤΙΜΩΝ:
       * Πίνακας Α: Κύριες Τιμές Μονάδας & Μετακινήσεις [1.000€ Λεμεσός + 200€/10km + 8€/km Συνοδεία Αστυνομίας]
       * Πίνακας Β: Κουτάκια Κανονικού Σκυροδέματος [ROUND(Τιμή αγοράς + 15€ τοποθέτηση + Στήσιμο αντλίας / m³, 0)]
       * Πίνακας Γ: Αποκλειστικά το Sleeve της συγκεκριμένης δουλειάς [(Σωλήνας + Φύρα + 0,78€ Εργατικά + 0,0876€ Τέλλα) × 1,30]
       * Πίνακας Δ: Ζωντανός Έλεγχος Γενικού Συνόλου)
  4. Verification Gate: επανυπολογισμός σε απομονωμένο Excel (DispatchEx), έλεγχος για
     #VALUE!/#DIV/0!/#REF!/#NAME?, σύγκριση με ανεξάρτητο υπολογισμό Python.
"""
import sys
import os
import re
import json
import math
import shutil
import argparse

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.workbook.defined_name import DefinedName

ENGINE_VERSION = "v5.7"
RULES_DIR = os.path.dirname(os.path.abspath(__file__))
CATALOG_PATH = os.path.join(RULES_DIR, "ΤΙΜΟΚΑΤΑΛΟΓΟΣ_ΕΤΑΙΡΕΙΑΣ_ΠΑΣΣΑΛΩΝ.xlsx")
CATALOG_V2_PATH = os.path.join(RULES_DIR, "ΤΙΜΟΚΑΤΑΛΟΓΟΣ_ΕΤΑΙΡΕΙΑΣ_ΠΑΣΣΑΛΩΝ_v2.xlsx")
PROJECTS_ROOT = r"C:\Users\nikth\Desktop\NIKOLAS TRY\Projects\Piles"
S1 = "Οικονομική Προσφορά"
S2 = "Τεχνικοί Υπολογισμοί"
S3 = "Τιμές & Έλεγχος"

MISSING_MSG = "Δεν υπάρχει αυτό το είδος ({item}) στον τιμοκατάλογο πασσάλων. Καθοδήγησέ με με τιμή να ξέρω."


class MissingMaterialError(Exception):
    pass


class SpecError(Exception):
    pass


# =============================================================================
# 1. ΤΙΜΟΚΑΤΑΛΟΓΟΣ
# =============================================================================
CATALOG_KEYS = {
    "mob_base": (r"^Μετακίνηση", "Μετακίνηση γεωτρύπανου — βασική τιμή εντός Λεμεσού"),
    "mob_extra": (r"Επιβάρυνση.*10 km|Επιβάρυνση ανά επιπλέον", "Επιβάρυνση μετακίνησης ανά 10 km"),
    "police_km": (r"^Συνοδεία Αστυνομίας", "Συνοδεία Αστυνομίας ανά km"),
    "conc_place": (r"^Τοποθέτηση / Άντληση Κανονικού Σκυροδέματος", "Τοποθέτηση κανονικού σκυροδέματος"),
    "steel": (r"Σίδηρος κάθετων ράβδων", "Χάλυβας οπλισμού B500C"),
    "spiral_steel": (r"Σίδηρος σούστας", "Χάλυβας σούστας"),
    "breaking": (r"^Κεφαλοκοπή", "Κεφαλοκοπή πασσάλων"),
    "sl_pipe_len_32": (r"Sleeve Φ32.*Εμπορικό μήκος", "Εμπορικό μήκος σωλήνα Sleeve Φ32"),
    "sl_pipe_price_32": (r"Sleeve Φ32.*Τιμή σωλήνα", "Τιμή σωλήνα Sleeve Φ32"),
    "sl_pipe_len_38": (r"Sleeve Φ38.*Εμπορικό μήκος", "Εμπορικό μήκος σωλήνα Sleeve Φ38"),
    "sl_pipe_price_38": (r"Sleeve Φ38.*Τιμή σωλήνα", "Τιμή σωλήνα Sleeve Φ38"),
    "sl_pipe_len_50": (r"Sleeve Φ50.*Εμπορικό μήκος", "Εμπορικό μήκος σωλήνα Sleeve Φ50"),
    "sl_pipe_price_50": (r"Sleeve Φ50.*Τιμή σωλήνα", "Τιμή σωλήνα Sleeve Φ50"),
    "sl_labor": (r"Εργατικά τοποθέτησης Sleeve ανά σίδερο", "Εργατικά τοποθέτησης Sleeve ανά σίδερο"),
    "sl_tape": (r"Κόστος Τέλλας.*ανά σίδερο", "Κόστος Τέλλας ανά σίδερο"),
    "sl_markup": (r"Συντελεστής Overhead, Profit & Damages στα Sleeves", "Συντελεστής Overhead, Profit & Damages (×1,30)"),
    "c_density": (r"Πυκνότητα χάλυβα", "Πυκνότητα χάλυβα"),
    "c_loss": (r"απώλεια / υπερκατανάλωση", "Συντελεστής απώλειας σκυροδέματος"),
    "zone_km": (r"Βήμα χιλιομετρικής ζώνης|Βασική ακτίνα", "Βήμα χιλιομετρικής ζώνης (10 km)"),
    "c_barlen": (r"Μέγιστο εμπορικό μήκος", "Μέγιστο εμπορικό μήκος βέργας"),
    "c_lap": (r"μήκους μάτισης", "Συντελεστής μάτισης"),
    "c_overc": (r"Μέγιστο ύψος υπερκάλυψης", "Μέγιστο ύψος υπερκάλυψης σκυροδέματος"),
}


def load_catalog(path=CATALOG_PATH):
    if os.path.exists(CATALOG_V2_PATH):
        try:
            shutil.copyfile(CATALOG_V2_PATH, CATALOG_PATH)
        except PermissionError:
            pass
        use_path = CATALOG_V2_PATH
    else:
        use_path = path
    if not os.path.exists(use_path):
        raise SpecError(f"Δεν βρέθηκε ο τιμοκατάλογος: {use_path}")
    wb = openpyxl.load_workbook(use_path, data_only=True)
    ws = wb["1_Τιμές_Εταιρείας"]
    items = []
    for r in range(1, ws.max_row + 1):
        desc, val, unit = ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value
        if isinstance(desc, str) and isinstance(val, (int, float)) and not isinstance(val, bool):
            items.append({"row": r, "desc": desc.strip(), "price": float(val), "unit": unit})
    return items


def catalog_find(items, pattern):
    hits = [it for it in items if re.search(pattern, it["desc"])]
    if len(hits) > 1:
        raise SpecError(f"Ο τιμοκατάλογος έχει >1 γραμμές για '{pattern}': " + ", ".join(h["desc"] for h in hits))
    return hits[0] if hits else None


class PriceBook:
    """Επιλύει τιμές: πρώτα ρητή τιμή χρήστη, μετά τιμοκατάλογος. Συλλέγει ό,τι λείπει."""

    def __init__(self, items, user_prices):
        self.items = items
        self.user = {k: float(v) for k, v in (user_prices or {}).items()}
        self.missing = []
        self.source = {}

    def get(self, key, pattern, label, fallback_val=None, fallback_src=None):
        if key in self.user:
            self.source[key] = "Τιμή χρήστη (εκτός τιμοκαταλόγου)"
            return self.user[key]
        hit = catalog_find(self.items, pattern)
        if hit is None:
            if fallback_val is not None:
                self.source[key] = fallback_src or "Προεπιλογή κανόνα"
                return fallback_val
            self.missing.append(label)
            self.source[key] = "ΛΕΙΠΕΙ"
            return None
        self.source[key] = f"Τιμοκατάλογος (γρ. {hit['row']})"
        return hit["price"]

    def raise_if_missing(self):
        if self.missing:
            msgs = [MISSING_MSG.format(item=m) for m in self.missing]
            raise MissingMaterialError("\n".join(msgs))


def select_project_sleeve_params(y_vert_mm, prices, pb_source):
    """
    Επιλέγει ΜΟΝΟ το Sleeve της συγκεκριμένης δουλειάς βάσει της διαμέτρου Φ της κάθετης ράβδου:
      - Έως Φ20      -> Sleeve Φ32 (σωλήνας 3.00m @ 4.0200€ = 1.3400 €/m)
      - Φ22 έως Φ25  -> Sleeve Φ38 (σωλήνας 3.00m @ 5.7240€ = 1.9080 €/m)
      - Φ26 έως Φ32  -> Sleeve Φ50 (σωλήνας 6.00m @ 12.0615€ = 2.01025 €/m)
    """
    y_int = int(round(y_vert_mm))
    if y_int <= 20:
        return {
            "type_label": f"Sleeve Φ32 (για κάθετο σίδερο Φ{y_int} ≤ Φ20)",
            "short_type": "Sleeve Φ32",
            "pipe_len": prices["sl_pipe_len_32"],
            "pipe_price": prices["sl_pipe_price_32"],
            "src_len": pb_source.get("sl_pipe_len_32", "Τιμοκατάλογος (Φύλλο 5)"),
            "src_price": pb_source.get("sl_pipe_price_32", "Τιμοκατάλογος (Φύλλο 5)"),
        }
    elif y_int <= 25:
        return {
            "type_label": f"Sleeve Φ38 (για κάθετο σίδερο Φ{y_int} [Φ22–Φ25])",
            "short_type": "Sleeve Φ38",
            "pipe_len": prices["sl_pipe_len_38"],
            "pipe_price": prices["sl_pipe_price_38"],
            "src_len": pb_source.get("sl_pipe_len_38", "Τιμοκατάλογος (Φύλλο 5)"),
            "src_price": pb_source.get("sl_pipe_price_38", "Τιμοκατάλογος (Φύλλο 5)"),
        }
    else:
        return {
            "type_label": f"Sleeve Φ50 (για κάθετο σίδερο Φ{y_int} [Φ26–Φ32])",
            "short_type": "Sleeve Φ50",
            "pipe_len": prices["sl_pipe_len_50"],
            "pipe_price": prices["sl_pipe_price_50"],
            "src_len": pb_source.get("sl_pipe_len_50", "Τιμοκατάλογος (Φύλλο 5)"),
            "src_price": pb_source.get("sl_pipe_price_50", "Τιμοκατάλογος (Φύλλο 5)"),
        }


def calc_billed_sleeve_len(req_len, pipe_len):
    if req_len <= 0:
        return 0.0
    if req_len <= pipe_len:
        pieces = int(math.floor(round(pipe_len / req_len, 6)))
        return pipe_len / max(1, pieces)
    return math.ceil(round(req_len / pipe_len, 6)) * pipe_len


# =============================================================================
# 2. CELL REGISTRY
# =============================================================================
TOKEN = re.compile(r"\{([A-Za-z0-9_.]+)\}")
RAW_ADDR = re.compile(r"(?<![A-Za-z_])\$?[A-Z]{1,3}\$?[0-9]+(?![0-9A-Za-z_(])")
STR_LITERAL = re.compile(r'"[^"]*"')


class Book:
    def __init__(self):
        self.wb = openpyxl.Workbook()
        self.ws1 = self.wb.active
        self.ws1.title = S1
        self.ws2 = self.wb.create_sheet(S2)
        self.ws3 = self.wb.create_sheet(S3)
        self.names = {}
        self.pending = []

    def reg(self, name, ws, coord):
        if name in self.names:
            raise SpecError(f"Διπλό όνομα κελιού: {name}")
        self.names[name] = (ws.title, coord)

    def put(self, ws, coord, value, name=None, fmt=None, font=None, fill=None, border=None, align=None):
        cell = ws[coord]
        if isinstance(value, str) and value.startswith("="):
            no_strings = STR_LITERAL.sub("", value)
            stripped = TOKEN.sub("", no_strings)
            if RAW_ADDR.search(stripped):
                raise SpecError(f"Καρφωτή διεύθυνση κελιού στη φόρμουλα '{value}' ({ws.title}!{coord}). "
                                f"Χρησιμοποίησε {{όνομα}}.")
            self.pending.append((ws, coord, value))
        else:
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
        if name:
            self.reg(name, ws, coord)
        if fmt:
            cell.number_format = fmt
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if border:
            cell.border = border
        if align:
            cell.alignment = align
        return cell

    def ref(self, name, from_ws):
        sh, coord = self.names[name]
        col, row = coordinate_from_string(coord)
        a = f"${col}${row}"
        return a if sh == from_ws.title else f"'{sh}'!{a}"

    def finalize(self):
        for ws, coord, tpl in self.pending:
            def sub(m):
                n = m.group(1)
                if n not in self.names:
                    raise SpecError(f"Άγνωστο όνομα κελιού {{{n}}} στη φόρμουλα του {ws.title}!{coord}")
                return self.ref(n, ws)
            ws[coord].value = TOKEN.sub(sub, tpl)
        for n, (sh, coord) in self.names.items():
            col, row = coordinate_from_string(coord)
            dn_name = "N_" + n.replace(".", "_")
            dn = DefinedName(dn_name, attr_text=f"'{sh}'!${col}${row}")
            self.wb.defined_names[dn_name] = dn

    def layout(self):
        return {n: {"sheet": sh, "cell": c} for n, (sh, c) in self.names.items()}


# =============================================================================
# 3. ΣΤΥΛ
# =============================================================================
def side(c, s="thin"):
    return Side(style=s, color=c)


ST = {
    "title_font": Font(name="Calibri", size=13, bold=True, color="FFFFFF"),
    "title_fill": PatternFill("solid", start_color="1E3A8A"),
    "sub_font": Font(name="Calibri", size=11, bold=True, color="1E40AF"),
    "meta_font": Font(name="Calibri", size=10, italic=True, color="4B5563"),
    "hdr_font": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
    "hdr_fill": PatternFill("solid", start_color="1E40AF"),
    "yellow": PatternFill("solid", start_color="FEF9C3"),
    "yellow_border": Border(left=side("D97706", "medium"), right=side("D97706", "medium"),
                            top=side("D97706", "medium"), bottom=side("D97706", "medium")),
    "linked_fill": PatternFill("solid", start_color="F1F5F9"),
    "gold": PatternFill("solid", start_color="FEF3C7"),
    "sub_row_fill": PatternFill("solid", start_color="F8FAFC"),
    "sub_row_font": Font(name="Calibri", size=9, italic=True, color="334155"),
    "red_fill": PatternFill("solid", start_color="FEE2E2"),
    "red_font": Font(name="Calibri", size=10, bold=True, color="991B1B"),
    "red_border": Border(left=side("DC2626", "medium"), right=side("DC2626", "medium"),
                         top=side("DC2626", "medium"), bottom=side("DC2626", "medium")),
    "em_fill": PatternFill("solid", start_color="D1FAE5"),
    "em_font": Font(name="Calibri", size=10, bold=True, color="065F46"),
    "em_title": Font(name="Calibri", size=11, bold=True, color="065F46"),
    "em_border": Border(left=side("10B981"), right=side("10B981"), top=side("10B981"), bottom=side("10B981")),
    "thin": Border(left=side("D1D5DB"), right=side("D1D5DB"), top=side("D1D5DB"), bottom=side("D1D5DB")),
    "gt_border": Border(left=side("B45309"), right=side("B45309"), top=side("B45309"), bottom=side("B45309", "double")),
    "gt_font": Font(name="Calibri", size=11, bold=True, color="92400E"),
    "note_font": Font(name="Calibri", size=9, color="374151"),
    "bold": Font(name="Calibri", size=10, bold=True),
}
FMT = {
    "m": "0.00", "m³": "0.00", "kg": "#,##0.0", "kg/m": "0.000", "mm": "0", "τεμ.": "#,##0",
    "πάσσαλο": "#,##0", "τρύπες": "#,##0", "€": '#,##0.00 "€"', "€0": '#,##0 "€"', "€4": '#,##0.0000 "€"',
    "€/m": '#,##0.0000 "€/m"', "€/km": '#,##0.00 "€/km"', "×": "0.00",
    "km": "0.0", "%": "0.0%", "ζώνες": "0", "-": "0.00000",
}
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


# =============================================================================
# 4. ΟΡΙΣΜΟΙ ΔΕΔΟΜΕΝΩΝ ΟΜΑΔΩΝ & ΒΗΜΑΤΩΝ ΥΠΟΛΟΓΙΣΜΟΥ (ΑΝΘΡΩΠΙΝΑ ΕΛΛΗΝΙΚΑ)
# =============================================================================
GROUP_INPUTS = [
    ("count", "Πλήθος πασσάλων ομάδας", "τεμ.", True, None),
    ("D_mm", "Διάμετρος τρύπας πασσάλου", "mm", True, None),
    ("L_pile", "Ύψος σκυροδετημένου πασσάλου (κάτω από το Cut-Off)", "m", True, None),
    ("cutoff", "Στάθμη αποκοπής κεφαλής πασσάλου (Cut-Off)", "m", False, None),
    ("raft", "Βάθος / Πάχος πεδιλόπλακας", "m", False, None),
    ("n_vert", "Κάθετα σίδερα κλωβού — Πλήθος", "τεμ.", True, None),
    ("y_vert", "Κάθετα σίδερα κλωβού — Διάμετρος Φ", "mm", True, None),
    ("starter", "Ύψος αναμονής / Starter (κενό σχεδίου → πεδιλόπλακα − 0.10m)", "m", False, None),
    ("top_bend", "Γωνία πάνω (κενό σχεδίου → 0.50m)", "m", False, None),
    ("bottom_bend", "Γωνία κάτω (στον πυθμένα)", "m", False, 0.0),
    ("lap_len", "Ύψος μάτισης (κενό σχεδίου → 50×Φ)", "m", False, None),
    ("y_spiral", "Σούστα & στεφάνια — Διάμετρος Φ", "mm", True, None),
    ("s_end", "Βήμα σούστας στα άκρα (πρώτα & τελευταία μέτρα)", "mm", False, None),
    ("end_zone", "Μήκος ζώνης πύκνωσης άκρου (πάνω & κάτω)", "m", False, 0.0),
    ("s_mid", "Βήμα σούστας (απόσταση μεταξύ σπειρών)", "mm", True, None),
    ("head_links", "Στεφάνια αναμονής starter — ρητά από σχέδιο (κενό → υπολογισμός)", "τεμ.", False, None),
    ("extra_n", "Πρόσθετος οπλισμός (U-bars κ.λπ.) — Πλήθος", "τεμ.", False, 0),
    ("extra_y", "Πρόσθετος οπλισμός — Διάμετρος Φ", "mm", False, 0),
    ("extra_len", "Πρόσθετος οπλισμός — Μήκος ανά βέργα", "m", False, 0.0),
]

MATH_STEPS = [
    ("dh", "Κενό τρύπημα (από πλατφόρμα μέχρι Cut-Off)", "m",
     "={wpl}-{G.cutoff}", False,
     "Κενό τρύπημα = Στάθμη πλατφόρμας − Στάθμη αποκοπής (Cut-Off)"),
    ("h_overc", "Ύψος υπερκάλυψης κουγκριού (πάνω από το Cut-Off)", "m",
     "=IF({G.dh}>0,MIN({c_overc},{G.dh}),0)", False,
     "Υπερκάλυψη κουγκριού = έως +2.00m πάνω από το Cut-Off (όταν υπάρχει κενό τρύπημα)"),
    ("L_drill", "Συνολικό βάθος τρυπήματος ανά πάσσαλο", "m",
     "={G.L_pile}+MAX(0,{G.dh})", True,
     "Συνολικό βάθος τρυπήματος = Ύψος σκυροδετημένου πασσάλου + Κενό τρύπημα"),
    ("V_exc", "Γεωμετρικός όγκος εκσκαφής τρύπας", "m³",
     "={c_pi}*({G.D_mm}/1000)^2/4*{G.L_drill}", True,
     "Όγκος εκσκαφής = Εμβαδόν τρύπας (π × Διάμετρος² / 4) × Συνολικό βάθος τρυπήματος"),
    ("L_conc", "Συνολικό ύψος σκυροδέτησης (με την υπερκάλυψη)", "m",
     "={G.L_pile}+{G.h_overc}", False,
     "Ύψος σκυροδέτησης = Ύψος σκυροδετημένου πασσάλου + Υπερκάλυψη κουγκριού"),
    ("V_conc", "Κυβικά σκυροδέματος ανά πάσσαλο (με +10% απώλεια)", "m³",
     "={c_pi}*({G.D_mm}/1000)^2/4*{G.L_conc}*(1+{c_loss})", True,
     "Κυβικά κουγκριού = Εμβαδόν τρύπας × Ύψος σκυροδέτησης × 1.10 (+10% απώλεια)"),
    ("D_cage", "Διάμετρος στεφανιού κλωβού (έξω-έξω)", "mm",
     "={G.D_mm}-2*{cover_mm}", False,
     "Διάμετρος κλωβού = Διάμετρος τρύπας − (2 × Επικάλυψη Cover)"),
    ("P_cage", "Περίμετρος στεφανιού κλωβού (μήκος 1 σπείρας)", "m",
     "={c_pi}*{G.D_cage}/1000", False,
     "Περίμετρος στεφανιού = π (3.1416) × Διάμετρος στεφανιού κλωβού"),
    ("L_base", "Βασικό μήκος κάθετης βέργας (χωρίς μάτιση)", "m",
     "={G.L_pile}+{G.starter}+{G.top_bend}+{G.bottom_bend}", False,
     "Βασικό μήκος βέργας = Ύψος πασσάλου + Ύψος αναμονής starter + Γωνία πάνω + Γωνία κάτω"),
    ("n_laps", "Αριθμός ματίσεων ανά κάθετη βέργα", "τεμ.",
     "=IF({G.L_base}>{c_barlen},INT(({G.L_base}-0.001)/{c_barlen}),0)", False,
     "Αριθμός ματίσεων = 1 μάτιση για κάθε 12 μέτρα βέργας (όταν ξεπερνά τα 12.00m)"),
    ("L_bar", "Συνολικό μήκος κάθετης βέργας (με τη μάτιση)", "m",
     "={G.L_base}+{G.n_laps}*{G.lap_len}", False,
     "Συνολικό μήκος βέργας = Βασικό μήκος βέργας + (Αριθμός ματίσεων × Ύψος μάτισης)"),
    ("w_vert", "Βάρος ενός μέτρου κάθετης βέργας", "kg/m",
     "={c_pi}*({G.y_vert}/1000)^2/4*{c_density}", False,
     "Βάρος μέτρου κάθετου σιδήρου = Εμβαδόν διατομής σιδήρου × 7,850 kg/m³"),
    ("W_vert", "Βάρος κάθετων σιδήρων ανά πάσσαλο", "kg",
     "={G.n_vert}*{G.L_bar}*{G.w_vert}", True,
     "Βάρος κάθετων σιδήρων = Πλήθος κάθετων σιδήρων × Συνολικό μήκος βέργας × Βάρος μέτρου"),
    ("n_turns", "Αριθμός σπειρών σούστας στο σώμα πασσάλου (κάτω από Cut-Off, με +1)", "τεμ.",
     "=ROUNDUP(ROUND(MIN({G.L_pile},2*{G.end_zone})*1000/{G.s_end}"
     "+MAX(0,{G.L_pile}-2*{G.end_zone})*1000/{G.s_mid},6),0)+1", False,
     "Αριθμός σπειρών σούστας = ROUNDUP(Ύψος πασσάλου ÷ Βήμα σούστας) + 1 (κάτω από το Cut-Off μπαίνει +1)"),
    ("w_spiral", "Βάρος ενός μέτρου σιδήρου σούστας", "kg/m",
     "={c_pi}*({G.y_spiral}/1000)^2/4*{c_density}", False,
     "Βάρος μέτρου σούστας = Εμβαδόν διατομής σούστας × 7,850 kg/m³"),
    ("W_spiral", "Βάρος σούστας σώματος ανά πάσσαλο", "kg",
     "={G.n_turns}*{G.P_cage}*{G.w_spiral}", True,
     "Βάρος σούστας σώματος = Αριθμός σπειρών × Περίμετρος στεφανιού × Βάρος μέτρου"),
    ("n_head", "Αριθμός στεφανιών στην αναμονή starter (πάνω από Cut-Off, χωρίς +1)", "τεμ.",
     "=ROUNDUP(ROUND({G.starter}*1000/IF({G.end_zone}>0,{G.s_end},{G.s_mid}),6),0)", False,
     "Στεφάνια αναμονής starter = ROUNDUP(Ύψος αναμονής ÷ Βήμα) — χωρίς +1"),
    ("W_head", "Βάρος στεφανιών αναμονής starter ανά πάσσαλο", "kg",
     "={G.n_head}*{G.P_cage}*{G.w_spiral}", True,
     "Βάρος στεφανιών αναμονής = Αριθμός στεφανιών αναμονής × Περίμετρος στεφανιού × Βάρος μέτρου"),
    ("W_extra", "Βάρος πρόσθετου οπλισμού (U-bars κ.λπ.) ανά πάσσαλο", "kg",
     "={G.extra_n}*{G.extra_len}*{c_pi}*({G.extra_y}/1000)^2/4*{c_density}", True,
     "Βάρος πρόσθετου οπλισμού = Πλήθος × Μήκος βέργας × Βάρος μέτρου"),
    ("W_total", "Συνολικό βάρος σιδήρου κλωβού ανά πάσσαλο", "kg",
     "={G.W_vert}+{G.W_spiral}+{G.W_head}+{G.W_extra}", True,
     "Συνολικό βάρος σιδήρου = Κάθετα σίδερα + Σούστα σώματος + Στεφάνια αναμονής + Πρόσθετος"),
    ("sleeves", "Προστατευτικά sleeves αναμονών ανά πάσσαλο", "τεμ.",
     "=IF({G.dh}>0,{G.n_vert},0)", True,
     "Sleeves ανά πάσσαλο = 1 τεμάχιο σε κάθε κάθετο σίδερο (κοστολόγηση στο Φύλλο 3 με ×1,30)"),
    ("breaking", "Κεφαλοκοπή (σπάσιμο κεφαλής) ανά πάσσαλο", "τεμ.",
     "=IF({G.dh}>0,1,0)", True,
     "Κεφαλοκοπή ανά πάσσαλο = 1 τεμάχιο ανά πάσσαλο (όταν υπάρχει κενό τρύπημα)"),
]

RULE_STARTER_OFFSET = 0.10
RULE_TOP_BEND = 0.50


def normalize_spec(spec):
    if not spec.get("groups"):
        raise SpecError("Το spec δεν έχει ομάδες πασσάλων (groups).")
    default_cut = (spec.get("elevations") or {}).get("cutoff")
    out = []
    for i, g in enumerate(spec["groups"], 1):
        ng = {"label": g.get("label", f"Ομάδα {i}"), "ids": g.get("ids", "")}
        if "pitch" in g and "s_mid" not in g:
            g = dict(g, s_mid=g["pitch"])
        if g.get("cutoff") is None and default_cut is not None:
            g = dict(g, cutoff=default_cut)
        for key, desc, unit, req, default in GROUP_INPUTS:
            if key in g and g[key] is not None:
                ng[key] = float(g[key])
            elif req:
                raise SpecError(f"Ομάδα {i} ({ng['label']}): λείπει το υποχρεωτικό πεδίο '{key}' ({desc}).")
            else:
                ng[key] = default
        if ng["s_end"] is None:
            ng["s_end"] = ng["s_mid"]
        if ng["starter"] is None and ng["raft"] is None:
            raise SpecError(f"Ομάδα {i} ({ng['label']}): δώσε 'starter' από το σχέδιο ή 'raft' (βάθος πεδιλόπλακας) "
                            f"για τον κανόνα starter = πεδιλόπλακα − 0.10 m.")
        for key in ("count", "D_mm", "L_pile", "n_vert", "y_vert", "y_spiral", "s_mid", "s_end"):
            if ng[key] <= 0:
                raise SpecError(f"Ομάδα {i}: το '{key}' πρέπει να είναι > 0.")
        ng["count"] = int(round(ng["count"]))
        ng["D_mm"] = int(round(ng["D_mm"]))
        ng["starter_eff"] = ng["starter"] if ng["starter"] is not None else ng["raft"] - RULE_STARTER_OFFSET
        ng["top_bend_eff"] = ng["top_bend"] if ng["top_bend"] is not None else RULE_TOP_BEND
        out.append(ng)
    return out


# =============================================================================
# 5. ΑΝΕΞΑΡΤΗΤΟΣ ΥΠΟΛΟΓΙΣΜΟΣ PYTHON & ΖΩΝΤΑΝΕΣ ΑΝΘΡΩΠΙΝΕΣ ΕΠΕΞΗΓΗΣΕΙΣ (ΠΙΝΑΚΑΣ 2Β)
# =============================================================================
def python_reference(spec, groups, consts, prices, sl_info):
    el = spec.get("elevations") or {}
    wpl = el.get("wpl")
    cover = float(spec.get("cover_mm", 75))
    rho, loss, barlen, lap = consts["c_density"], consts["c_loss"], consts["c_barlen"], consts["c_lap"]

    def w(phi):
        return math.pi * (phi / 1000.0) ** 2 / 4.0 * rho

    n_piles = sum(g["count"] for g in groups)
    drill_by_d, conc, steel, sleeves, breaking = {}, 0.0, 0.0, 0, 0
    tot_w_vert, tot_w_spiral, tot_w_head, tot_w_extra = 0.0, 0.0, 0.0, 0.0
    group_calcs = []
    for g in groups:
        dh = (float(wpl) - g["cutoff"]) if (wpl is not None and g["cutoff"] is not None) else 0.0
        hov = min(consts["c_overc"], dh) if dh > 0 else 0.0
        area = math.pi * (g["D_mm"] / 1000.0) ** 2 / 4.0
        l_drill = g["L_pile"] + max(0.0, dh)
        l_conc = g["L_pile"] + hov
        v_conc = area * l_conc * (1 + loss)
        starter = g["starter_eff"]
        l_base = g["L_pile"] + starter + g["top_bend_eff"] + g["bottom_bend"]
        laps = math.floor((l_base - 0.001) / barlen) if l_base > barlen else 0
        lap_len_eff = g["lap_len"] if g.get("lap_len") is not None else (lap * g["y_vert"] / 1000.0)
        l_bar = l_base + laps * lap_len_eff
        w_v = w(g["y_vert"])
        w_vert = g["n_vert"] * l_bar * w_v
        ends = min(g["L_pile"], 2 * g["end_zone"])
        mid = max(0.0, g["L_pile"] - 2 * g["end_zone"])
        turns = math.ceil(round(ends * 1000 / g["s_end"] + mid * 1000 / g["s_mid"], 6)) + 1
        p_cage = math.pi * (g["D_mm"] - 2 * cover) / 1000.0
        w_sp = w(g["y_spiral"])
        w_spiral = turns * p_cage * w_sp
        s_top = g["s_end"] if g["end_zone"] > 0 else g["s_mid"]
        n_head = int(round(g["head_links"])) if g["head_links"] is not None else math.ceil(round(starter * 1000 / s_top, 6))
        w_head = n_head * p_cage * w_sp
        w_extra = g["extra_n"] * g["extra_len"] * w(g["extra_y"]) if g["extra_y"] > 0 else 0.0
        n = g["count"]
        drill_by_d[g["D_mm"]] = drill_by_d.get(g["D_mm"], 0.0) + n * l_drill
        conc += n * v_conc
        tot_w_vert += n * w_vert
        tot_w_spiral += n * w_spiral
        tot_w_head += n * w_head
        tot_w_extra += n * w_extra
        steel += n * (w_vert + w_spiral + w_head + w_extra)
        sleeves += n * (int(round(g["n_vert"])) if dh > 0 else 0)
        breaking += n * (1 if dh > 0 else 0)
        group_calcs.append({
            "g": g, "dh": dh, "hov": hov, "area": area, "l_drill": l_drill, "l_conc": l_conc,
            "v_conc": v_conc, "starter": starter, "l_base": l_base, "laps": laps,
            "lap_len_eff": lap_len_eff, "l_bar": l_bar,
            "w_v": w_v, "w_vert": w_vert, "turns": turns, "p_cage": p_cage, "w_sp": w_sp,
            "w_spiral": w_spiral, "s_top": s_top, "n_head": n_head, "w_head": w_head,
            "w_extra": w_extra, "w_total": w_vert + w_spiral + w_head + w_extra,
        })

    # Υπολογισμός Μετακίνησης: 1.000€ εντός Λεμεσού + 200€ ανά 10km εκτός Λεμεσού + 8€/km Συνοδεία Αστυνομίας
    proj = spec.get("project", {})
    dist = float(proj.get("distance_km", 0))
    police_km_val = float(proj.get("police_km", dist))
    zones = math.ceil(round(dist / consts["zone_km"], 9)) if dist > 0 else 0
    mob_police = police_km_val * prices["police_km"]
    mob_total = prices["mob_base"] + zones * prices["mob_extra"] + mob_police

    # Υπολογισμός Τελικής Τιμής Κανονικού Σκυροδέματος (ΣΤΡΟΓΓΥΛΟΠΟΙΗΜΕΝΗΣ στο Τ.Β.7 / D21)
    avg_v_pile = conc / n_piles if n_piles > 0 else 0.0
    v_per_setup = prices["piles_per_setup"] * avg_v_pile
    pump_per_m3 = (prices["pump_setup"] / v_per_setup) if (prices["pump_setup"] > 0 and v_per_setup > 0) else 0.0
    p_conc_eff = float(math.floor(prices["conc_buy"] + prices["conc_place"] + pump_per_m3 + 0.5))

    # Υπολογισμός Τελικής Τιμής Sleeve της συγκεκριμένης δουλειάς (με × 1,30 Overhead, Profit & Damages)
    req_len = groups[0]["starter_eff"] + groups[0]["top_bend_eff"]
    billed_len = calc_billed_sleeve_len(req_len, sl_info["pipe_len"])
    price_per_m = sl_info["pipe_price"] / sl_info["pipe_len"]
    mat_cost = billed_len * price_per_m
    base_sl_cost = mat_cost + prices["sl_labor"] + prices["sl_tape"]
    p_sleeve_eff = base_sl_cost * prices["sl_markup"]

    lines = {"mob": n_piles * (mob_total / n_piles)}
    for d, m in drill_by_d.items():
        lines[f"drill_{d}"] = m * prices[f"drill_{d}"]
    lines["conc"] = conc * p_conc_eff
    lines["steel"] = steel * prices["steel"]
    lines["sleeve"] = sleeves * p_sleeve_eff
    lines["breaking"] = breaking * prices["breaking"]
    return {
        "wpl": wpl, "n_piles": n_piles, "zones": zones, "drill_m": drill_by_d, "conc_m3": conc,
        "steel_kg": steel, "tot_w_vert": tot_w_vert, "tot_w_spiral": tot_w_spiral,
        "tot_w_head": tot_w_head, "tot_w_extra": tot_w_extra, "sleeves": sleeves,
        "breaking": breaking, "mob_police": mob_police, "mob_total": mob_total,
        "p_conc_eff": p_conc_eff, "v_per_setup": v_per_setup, "pump_per_m3": pump_per_m3,
        "sl_base_cost": base_sl_cost, "p_sleeve_eff": p_sleeve_eff,
        "sl_req_len": req_len, "sl_billed_len": billed_len, "sl_mat_cost": mat_cost,
        "group_calcs": group_calcs, "lines": lines, "grand_total": sum(lines.values()),
    }


def build_consolidated_rows(spec, groups, diameters, consts, prices, cls, ref, sl_info):
    """
    Δημιουργεί τις γραμμές της Ενότητας 2Β (Συγκεντρωτικός Πίνακας Πράξεων, Ποσοτήτων & Τιμών).
    Όλες οι τιμές μονάδας αντλούνται από το Κεντρικό Φύλλο 3 («Τιμές & Έλεγχος») και όλες οι
    επεξηγήσεις είναι 100% ζωντανές φόρμουλες Excel σε καθαρά ανθρώπινα ελληνικά για τον μάστρο!
    """
    rows = []

    # 1. Μετακινήσεις (1.000€ Λεμεσός + 200€/10km + 8€/km Συνοδεία Αστυνομίας)
    mob_f = (
        '="Βασική Λεμεσού ("&ROUND({p_mob_base},2)&" €) + Εκτός Λεμεσού ανά "&ROUND({zone_km},0)&"km ("'
        '&{extra_zones}&" × "&ROUND({p_mob_extra},2)&" € = "&ROUND({extra_zones}*{p_mob_extra},2)'
        '&" €) + Συνοδεία Αστυνομίας ("&ROUND({police_km},1)&" km × "&ROUND({p_police_km},2)'
        '&" €/km = "&ROUND({mob_police},2)&" €) = Σύνολο ("&ROUND({mob_total},2)&" €) ÷ Πάσσαλοι ("'
        '&{n_piles}&") = "&ROUND({mob_per_pile},2)&" €/πάσσαλο"'
    )
    rows.append({
        "key": "mob", "code": "2Β.1",
        "desc": "Μετακινήσεις γεωτρύπανου & Συνοδεία Αστυνομίας (επιμερισμός ανά πάσσαλο)",
        "unit": "πάσσαλο", "f_qty": "={n_piles}", "f_price": "={mob_per_pile}",
        "is_main": True, "math": mob_f,
    })

    # 2. Εκσκαφή / Διάτρηση ανά διάμετρο
    for j, d in enumerate(diameters):
        sub = f"2Β.2{chr(ord('a') + j)}" if len(diameters) > 1 else "2Β.2"
        idx = [i for i, g in enumerate(groups) if g["D_mm"] == d]
        cnt = sum(groups[i]["count"] for i in idx)
        g_parts = [
            f'{{g{i}.count}}&" πασσ. × [Ύψος πασσάλου ("&ROUND({{g{i}.L_pile}},3)&"m) + Κενό τρύπημα ("'
            f'&ROUND({{g{i}.dh}},3)&"m) = "&ROUND({{g{i}.L_drill}},2)&"m]"'
            for i in idx
        ]
        drill_f = (
            '="Τρύπημα: "&' + ' & " + " & '.join(g_parts)
            + f'&" = Σύνολο ("&ROUND({{b2.drill_{d}.qty}},2)&"m) × Τιμή μέτρου ("&ROUND({{b2.drill_{d}.price}},2)'
            + f'&" €) = "&ROUND({{b2.drill_{d}.total}},2)&" €"'
        )
        rows.append({
            "key": f"drill_{d}", "code": sub,
            "desc": f"Εκσκαφή / Διάτρηση πασσάλων Ø{d}mm ({cnt} τεμ.)",
            "unit": "m", "f_qty": f"={{q.drill_{d}}}", "f_price": f"={{p.drill_{d}}}",
            "is_main": True, "math": drill_f,
        })

    # 3. Σκυρόδεμα (με ανάλυση: ROUND(Αγορά + 15€ Τοποθέτηση + Στήσιμο Αντλίας/m³, 0))
    conc_f = (
        '="Κουγκρί ανά πάσσαλο: Εμβαδόν Ø"&{g0.D_mm}&"mm × [Ύψος πασσάλου ("&ROUND({g0.L_pile},3)'
        '&"m) + Υπερκάλυψη ("&ROUND({g0.h_overc},2)&"m) = "&ROUND({g0.L_conc},3)&"m] × Απώλεια (+"'
        '&ROUND({c_loss}*100,0)&"%) = "&ROUND({g0.V_conc},3)&" m³/πασσ. → Σύνολο ("&ROUND({b2.conc.qty},2)'
        '&" m³) × Στρογγυλοποιημένη Τιμή κυβικού [ROUND(Αγορά "&ROUND({p.conc_buy},2)&" € + Τοποθέτηση "'
        '&ROUND({p.conc_place},2)&" € + Αντλία "&ROUND({p.pump_per_m3},2)&" €, 0) = "'
        '&ROUND({b2.conc.price},0)&" €/m³] = "&ROUND({b2.conc.total},2)&" €"'
    )
    rows.append({
        "key": "conc", "code": "2Β.3",
        "desc": '="Σκυρόδεμα πασσάλων "&{conc_class}&" (με +"&ROUND({c_loss}*100,0)&"% απώλεια)"',
        "unit": "m³", "f_qty": "={q.conc}", "f_price": "={p.conc}",
        "is_main": True, "math": conc_f,
    })

    # 4α. Κάθετα σίδερα κλωβού
    vert_f = (
        '="Μήκος κάθετης βέργας = Ύψος πασσάλου ("&ROUND({g0.L_pile},3)&"m) + Ύψος αναμονής starter ("'
        '&ROUND({g0.starter},2)&"m) + Γωνία πάνω ("&ROUND({g0.top_bend},2)&"m) + Γωνία κάτω ("'
        '&ROUND({g0.bottom_bend},2)&"m) + Ύψος μάτισης ("&{g0.n_laps}&" μάτιση × "&ROUND({g0.lap_len},2)&"m) = "'
        '&ROUND({g0.L_bar},3)&"m → "&{g0.n_vert}&" σίδερα Φ"&{g0.y_vert}&" × "&ROUND({g0.L_bar},3)'
        '&"m × Βάρος μέτρου ("&ROUND({g0.w_vert},3)&" kg/m) = "&ROUND({g0.W_vert},2)&" kg/πασσ. → Σύνολο ("'
        '&ROUND({b2.steel_vert.qty},1)&" kg) × "&ROUND({b2.steel_vert.price},2)&" € = "'
        '&ROUND({b2.steel_vert.total},2)&" €"'
    )
    rows.append({
        "key": "steel_vert", "code": "2Β.4α",
        "desc": "   ↳ Οπλισμός: Κάθετα σίδερα κλωβού",
        "unit": "kg", "f_qty": "={t.W_vert}", "f_price": "={p.steel}",
        "is_main": False, "math": vert_f,
    })

    # 4β. Σούστα σώματος πασσάλου (κάτω από Cut-Off, με +1)
    sp_f = (
        '="Αριθμός σπειρών σούστας = [Ύψος πασσάλου ("&ROUND({g0.L_pile},3)&"m) ÷ Βήμα σούστας ("'
        '&ROUND({g0.s_mid}/1000,3)&"m) = "&({g0.n_turns}-1)&"] + 1 = "&{g0.n_turns}'
        '&" σπείρες × Περίμετρος στεφανιού κλωβού ("&ROUND({g0.P_cage},3)&"m) × Βάρος μέτρου Φ"&{g0.y_spiral}'
        '&" ("&ROUND({g0.w_spiral},3)&" kg/m) = "&ROUND({g0.W_spiral},2)&" kg/πασσ. → Σύνολο ("'
        '&ROUND({b2.steel_spiral.qty},1)&" kg) × "&ROUND({b2.steel_spiral.price},2)&" € = "'
        '&ROUND({b2.steel_spiral.total},2)&" €"'
    )
    rows.append({
        "key": "steel_spiral", "code": "2Β.4β",
        "desc": "   ↳ Οπλισμός: Σούστα σώματος πασσάλου (κάτω από το Cut-Off, με +1)",
        "unit": "kg", "f_qty": "={t.W_spiral}", "f_price": "={p.steel}",
        "is_main": False, "math": sp_f,
    })

    # 4γ. Στεφάνια αναμονής starter (πάνω από Cut-Off, χωρίς +1)
    hd_f = (
        '="Αριθμός στεφανιών αναμονής = Ύψος αναμονής starter ("&ROUND({g0.starter},2)&"m) ÷ Βήμα ("'
        '&ROUND({g0.s_mid}/1000,3)&"m) = "&{g0.n_head}&" στεφάνια (χωρίς +1) × Περίμετρος στεφανιού κλωβού ("'
        '&ROUND({g0.P_cage},3)&"m) × Βάρος μέτρου Φ"&{g0.y_spiral}&" ("&ROUND({g0.w_spiral},3)&" kg/m) = "'
        '&ROUND({g0.W_head},2)&" kg/πασσ. → Σύνολο ("&ROUND({b2.steel_head.qty},1)&" kg) × "'
        '&ROUND({b2.steel_head.price},2)&" € = "&ROUND({b2.steel_head.total},2)&" €"'
    )
    rows.append({
        "key": "steel_head", "code": "2Β.4γ",
        "desc": "   ↳ Οπλισμός: Στεφάνια αναμονής starter (πάνω από το Cut-Off, χωρίς +1)",
        "unit": "kg", "f_qty": "={t.W_head}", "f_price": "={p.steel}",
        "is_main": False, "math": hd_f,
    })

    # 4δ. Πρόσθετος οπλισμός (αν υπάρχει)
    if ref["tot_w_extra"] > 0:
        ex_f = (
            '="Πρόσθετος οπλισμός: Σύνολο ("&ROUND({b2.steel_extra.qty},1)&" kg) × Τιμή κιλού ("'
            '&ROUND({b2.steel_extra.price},2)&" €) = "&ROUND({b2.steel_extra.total},2)&" €"'
        )
        rows.append({
            "key": "steel_extra", "code": "2Β.4δ",
            "desc": "   ↳ Οπλισμός: Πρόσθετος οπλισμός (U-bars κ.λπ.)",
            "unit": "kg", "f_qty": "={t.W_extra}", "f_price": "={p.steel}",
            "is_main": False, "math": ex_f,
        })

    # 4. Σύνολο Χάλυβα
    st_f = (
        '="Συνολικό βάρος ανά πάσσαλο = Κάθετα σίδερα ("&ROUND({g0.W_vert},2)&" kg) + Σούστα σώματος ("'
        '&ROUND({g0.W_spiral},2)&" kg) + Στεφάνια αναμονής ("&ROUND({g0.W_head},2)&" kg) = "'
        '&ROUND({g0.W_total},2)&" kg/πασσ. → Σύνολο έργου ("&ROUND({b2.steel.qty},1)&" kg) × Τιμή κιλού ("'
        '&ROUND({b2.steel.price},2)&" €) = "&ROUND({b2.steel.total},2)&" € ("'
        '&ROUND({b2.steel.total}/{n_piles},2)&" €/πασσ.)"'
    )
    rows.append({
        "key": "steel", "code": "2Β.4",
        "desc": "Χάλυβας οπλισμού B500C — ΣΥΝΟΛΟ ΚΛΩΒΟΥ (4α + 4β + 4γ)",
        "unit": "kg", "f_qty": "={q.steel}", "f_price": "={p.steel}",
        "is_main": True, "math": st_f,
    })

    # 5. Sleeves: [(Υλικό σωλήνα + Εργατικά + Τέλλα) × 1,30 Overhead/Profit/Damages]
    sl_f = (
        f'="{sl_info["short_type"]} (για σίδερα Φ"&{{g0.y_vert}}&"): [Χρεωστέο μήκος ("&ROUND({{sl_billed_len}},2)'
        f'&"m × "&ROUND({{sl_price_m}},4)&" €/m = "&ROUND({{sl_mat_cost}},4)&" €) + Εργατικά ("'
        f'&ROUND({{sl_labor}},4)&" €) + Τέλλα ("&ROUND({{sl_tape}},4)&" €) = "&ROUND({{sl_base_cost}},4)'
        f'&" €] × "&ROUND({{sl_markup}},2)&" (Overhead/Profit/Damages) = "&ROUND({{b2.sleeve.price}},4)'
        f'&" €/τεμ. × "&{{b2.sleeve.qty}}&" τεμ. = "&ROUND({{b2.sleeve.total}},2)&" €"'
    )
    rows.append({
        "key": "sleeve", "code": "2Β.5",
        "desc": f'="Sleeves προστασίας αναμονών starter ({sl_info["short_type"]} για κάθετα σίδερα Φ"&{{g0.y_vert}}&")"',
        "unit": "τεμ.", "f_qty": "={q.sleeve}", "f_price": "={p.sleeve}",
        "is_main": True, "math": sl_f,
    })

    # 6. Κεφαλοκοπή
    br_f = (
        '="Κεφαλοκοπή = "&{g0.breaking}&" σπάσιμο κεφαλής/πάσσαλο × "&{n_piles}&" πασσάλους = "'
        '&{b2.breaking.qty}&" τεμάχια × Τιμή τεμαχίου ("&ROUND({b2.breaking.price},2)&" €) = "'
        '&ROUND({b2.breaking.total},2)&" € ("&ROUND({b2.breaking.total}/{n_piles},2)&" €/πασσ.)"'
    )
    rows.append({
        "key": "breaking", "code": "2Β.6",
        "desc": "Κεφαλοκοπή πασσάλων (σπάσιμο κεφαλής / Pile Breaking)",
        "unit": "τεμ.", "f_qty": "={q.breaking}", "f_price": "={p.breaking}",
        "is_main": True, "math": br_f,
    })

    return rows


# =============================================================================
# 6. ΚΑΤΑΣΚΕΥΗ EXCEL (3 ΦΥΛΛΑ: 1. Προσφορά, 2. Υπολογισμοί, 3. Τιμές & Έλεγχος)
# =============================================================================
def build(spec, out_path):
    groups = normalize_spec(spec)
    items = load_catalog()
    pb = PriceBook(items, spec.get("user_prices"))

    consts = {}
    for key in ("c_density", "c_loss", "zone_km", "c_barlen", "c_lap", "c_overc"):
        pat, label = CATALOG_KEYS[key]
        consts[key] = pb.get(key, pat, label)

    prices = {}
    for key in ("mob_base", "mob_extra", "police_km", "conc_place", "steel", "spiral_steel", "breaking",
                "sl_pipe_len_32", "sl_pipe_price_32", "sl_pipe_len_38", "sl_pipe_price_38",
                "sl_pipe_len_50", "sl_pipe_price_50", "sl_labor", "sl_tape", "sl_markup"):
        pat, label = CATALOG_KEYS[key]
        prices[key] = pb.get(key, pat, label)

    diameters = sorted({g["D_mm"] for g in groups}, reverse=True)
    for d in diameters:
        prices[f"drill_{d}"] = pb.get(f"drill_{d}", rf"Τρύπημα / Διάτρηση Πασσάλου Ø{d}mm",
                                      f"Τρύπημα / Διάτρηση Πασσάλου Ø{d}mm")
    cls = spec.get("concrete_class")
    if not cls:
        raise SpecError("Λείπει η κατηγορία σκυροδέματος (concrete_class). Δήλωσε τι γράφει το σχέδιο.")

    # Κουτάκια Κανονικού Σκυροδέματος (Τιμή αγοράς + 15€ τοποθέτηση + Στήσιμο αντλίας / m³)
    user_p = spec.get("user_prices") or {}
    if "conc_buy" in user_p:
        prices["conc_buy"] = float(user_p["conc_buy"])
        pb.source["conc_buy"] = "Κουτάκι Μηχανικού (Τιμή αγοράς κουγκριού)"
    elif f"concrete_{cls}" in user_p:
        prices["conc_buy"] = float(user_p[f"concrete_{cls}"])
        pb.source["conc_buy"] = "Κουτάκι Μηχανικού (Τιμή αγοράς κουγκριού)"
    else:
        prices["conc_buy"] = pb.get("conc_buy", rf"Σκυρόδεμα {re.escape(cls)}\b", f"Τιμή αγοράς σκυροδέματος {cls}")

    prices["pump_setup"] = float(user_p.get("pump_setup", 0.0))
    pb.source["pump_setup"] = "Κουτάκι Μηχανικού (Κόστος στησίματος αντλίας)"
    prices["piles_per_setup"] = float(user_p.get("piles_per_setup", 6.0))
    pb.source["piles_per_setup"] = "Κουτάκι Μηχανικού (Τρύπες/πασσάλους ανά στήσιμο αντλίας)"

    pb.raise_if_missing()

    sl_info = select_project_sleeve_params(groups[0]["y_vert"], prices, pb.source)
    ref = python_reference(spec, groups, consts, {**prices}, sl_info)

    el = spec.get("elevations") or {}
    red_alert = el.get("wpl") is None or any(g["cutoff"] is None for g in groups)
    proj = spec.get("project", {})
    n = len(groups)

    b = Book()
    ws1, ws2, ws3 = b.ws1, b.ws2, b.ws3

    # =========================================================================
    # ΦΥΛΛΟ 3: «Τιμές & Έλεγχος» (Ο ΜΟΝΑΔΙΚΟΣ ΑΡΧΗΓΟΣ ΤΙΜΩΝ!)
    # =========================================================================
    ws3.merge_cells("A1:F1")
    b.put(ws3, "A1", "ΘΕΜΕΛΙΟΤΕΧΝΙΚΗ ΛΤΔ — ΚΕΝΤΡΙΚΟΣ ΠΙΝΑΚΑΣ ΤΙΜΩΝ & ΕΛΕΓΧΟΥ ΕΡΓΟΥ (ΑΛΛΑΖΕΙΣ ΤΙΜΕΣ ΜΟΝΟ ΕΔΩ!)",
          font=ST["title_font"], fill=ST["title_fill"], align=CENTER)
    ws3.row_dimensions[1].height = 28
    ws3.merge_cells("A2:F2")
    b.put(ws3, "A2", "Όλες οι τιμές μονάδας και τα κουτάκια μετακινήσεων/σκυροδέματος/sleeve ορίζονται ΑΠΟΚΛΕΙΣΤΙΚΑ εδώ και ενημερώνουν αυτόματα όλα τα φύλλα και τις σημειώσεις!",
          font=ST["meta_font"])

    r3 = 4
    b.put(ws3, f"A{r3}", "Α. ΚΥΡΙΕΣ ΤΙΜΕΣ ΜΟΝΑΔΑΣ & ΜΕΤΑΚΙΝΗΣΕΙΣ (ΚΙΤΡΙΝΟ = ΑΛΛΑΖΕΙΣ ΕΔΩ)", font=ST["sub_font"])
    r3 += 1
    for i, lab in enumerate(("Κωδ.", "Περιγραφή Τιμής / Παραμέτρου", "Μονάδα", "Τιμή [ΑΛΛΑΖΕΙΣ ΕΔΩ]", "Προέλευση", "Πού Εφαρμόζεται")):
        b.put(ws3, f"{get_column_letter(1 + i)}{r3}", lab, font=ST["hdr_font"], fill=ST["hdr_fill"], align=CENTER)
    ws3.row_dimensions[r3].height = 28
    r3 += 1

    def put_master_price(code, desc, unit, val, name, src, target_info, yellow=True, green=False, fmt_override=None):
        nonlocal r3
        b.put(ws3, f"A{r3}", code, align=CENTER, border=ST["thin"], font=ST["bold"])
        b.put(ws3, f"B{r3}", desc, align=LEFT_WRAP, border=ST["thin"], font=ST["bold"] if (yellow or green) else None)
        b.put(ws3, f"C{r3}", unit, align=CENTER, border=ST["thin"])
        cell_fill = ST["em_fill"] if green else (ST["yellow"] if yellow else ST["linked_fill"])
        cell_bdr = ST["em_border"] if green else (ST["yellow_border"] if yellow else ST["thin"])
        cell_font = ST["em_font"] if green else ST["bold"]
        c = b.put(ws3, f"D{r3}", val, name=name, fmt=fmt_override or FMT.get(unit, "#,##0.00"), align=RIGHT,
                  border=cell_bdr, fill=cell_fill, font=cell_font)
        b.put(ws3, f"E{r3}", src, align=LEFT_WRAP, border=ST["thin"], font=ST["note_font"])
        b.put(ws3, f"F{r3}", target_info, align=LEFT_WRAP, border=ST["thin"], font=ST["note_font"])
        ws3.row_dimensions[r3].height = 24
        r3 += 1
        return c

    dist_val = float(proj.get("distance_km", 0))
    put_master_price("Τ.1", "Χιλιόμετρα εκτός Λεμεσού (0 km αν το έργο είναι μέσα στη Λεμεσό)", "km",
                     dist_val, "dist_km",
                     proj.get("distance_note", "Οδηγία χρήστη"), "Υπολογίζει αυτόματα τις δεκάδες km (×200€) & τη Συνοδεία Αστυνομίας (×8€/km)")
    put_master_price("Τ.2", "Μετακίνηση γεωτρύπανου — Βασική τιμή (Μέσα στη Λεμεσό)", "€",
                     prices["mob_base"], "p_mob_base", pb.source["mob_base"], "Φύλλο 1 (Γρ. 1) & Φύλλο 2 (1.1.4 & 2Β.1)")
    put_master_price("Τ.3", "Μετακίνηση γεωτρύπανου — Επιβάρυνση ανά 10 km εκτός Λεμεσού (+200 €/10km)", "€",
                     prices["mob_extra"], "p_mob_extra", pb.source["mob_extra"], "Φύλλο 1 (Γρ. 1) & Φύλλο 2 (1.1.5 & 2Β.1)")
    for idx_d, d in enumerate(diameters, 1):
        put_master_price(f"Τ.4.{idx_d}", f"Τρύπημα / Διάτρηση Πασσάλων Ø{d}mm", "€",
                         prices[f"drill_{d}"], f"p.drill_{d}", pb.source[f"drill_{d}"],
                         "Φύλλο 1 (Γρ. 2) & Φύλλο 2 (2Β.2) — €/m")
    put_master_price("Τ.5", "Χάλυβας οπλισμού B500C (Κάθετα σίδερα & Σούστα) — Τιμή κιλού", "€",
                     prices["steel"], "p.steel", pb.source["steel"], "Φύλλο 1 (Γρ. 4) & Φύλλο 2 (2Β.4, 4α, 4β, 4γ) — €/kg")
    put_master_price("Τ.6", "Κεφαλοκοπή πασσάλων (Pile Breaking) — Τιμή τεμαχίου", "€",
                     prices["breaking"], "p.breaking", pb.source["breaking"], "Φύλλο 1 (Γρ. 6) & Φύλλο 2 (2Β.6) — €/τεμ.")

    # --- ΠΙΝΑΚΑΣ Β: ΚΟΥΤΑΚΙΑ ΚΑΝΟΝΙΚΟΥ ΣΚΥΡΟΔΕΜΑΤΟΣ (Το Τ.Β.7 παραμένει ακριβώς στο D21 και είναι ΣΤΡΟΓΓΥΛΟΠΟΙΗΜΕΝΟ!)
    r3 += 1
    b.put(ws3, f"A{r3}", f"Β. ΥΠΟΛΟΓΙΣΜΟΣ ΤΙΜΗΣ ΚΑΝΟΝΙΚΟΥ ΣΚΥΡΟΔΕΜΑΤΟΣ ({cls}) — ΚΟΥΤΑΚΙΑ ΜΗΧΑΝΙΚΟΥ", font=ST["sub_font"])
    r3 += 1
    for i, lab in enumerate(("Κωδ.", "Παράμετρος Κόστους Σκυροδέτησης", "Μονάδα", "Τιμή / Κουτάκι", "Προέλευση / Τύπος", "Επεξήγηση")):
        b.put(ws3, f"{get_column_letter(1 + i)}{r3}", lab, font=ST["hdr_font"], fill=ST["hdr_fill"], align=CENTER)
    ws3.row_dimensions[r3].height = 26
    r3 += 1

    put_master_price("Τ.Β.1", f"Τιμή αγοράς σκυροδέματος ({cls}) από προμηθευτή [ΚΟΥΤΑΚΙ ΜΗΧΑΝΙΚΟΥ]", "€",
                     prices["conc_buy"], "p.conc_buy", pb.source["conc_buy"],
                     "Συμπληρώνεται από τον μηχανικό ανάλογα με την προσφορά προμηθευτή (€/m³)", yellow=True)
    put_master_price("Τ.Β.2", "Κόστος τοποθέτησης σκυροδέματος (σταθερό ανά κυβικό)", "€",
                     prices["conc_place"], "p.conc_place", pb.source["conc_place"],
                     "Σταθερή χρέωση τοποθέτησης = +15,00 €/m³ από τον Τιμοκατάλογο Εταιρείας", yellow=True)
    put_master_price("Τ.Β.3", "Κόστος στησίματος αντλίας (πρέσας) ανά στήσιμο [ΚΟΥΤΑΚΙ ΜΗΧΑΝΙΚΟΥ]", "€",
                     prices["pump_setup"], "p.pump_setup", pb.source["pump_setup"],
                     "Συμπληρώνεται από τον μηχανικό (€/στήσιμο αντλίας)", yellow=True)
    put_master_price("Τ.Β.4", "Πόσες τρύπες (πασσάλους) θα κάνει η αντλία ανά στήσιμο [ΚΟΥΤΑΚΙ ΜΗΧΑΝΙΚΟΥ]", "τρύπες",
                     prices["piles_per_setup"], "p.piles_per_setup", pb.source["piles_per_setup"],
                     "Αριθμός πασσάλων που σκυροδετούνται σε 1 στήσιμο της αντλίας", yellow=True)
    put_master_price("Τ.Β.5", "Κυβικά κουγκριού που αντλεί η αντλία ανά στήσιμο (Τρύπες × m³/τρύπα)", "m³",
                     "={p.piles_per_setup}*({q.conc}/{n_piles})", "v_per_setup",
                     "Αυτόματος υπολογισμός: Τρύπες ανά στήσιμο × Μέσα κυβικά ανά πάσσαλο",
                     "Κυβικά σκυροδέματος ανά στήσιμο αντλίας", yellow=False)
    put_master_price("Τ.Β.6", "Επιβάρυνση στησίματος αντλίας ανά κυβικό (Στήσιμο ÷ m³ ανά στήσιμο)", "€",
                     "=IF(AND({p.pump_setup}>0,{v_per_setup}>0),{p.pump_setup}/{v_per_setup},0)", "p.pump_per_m3",
                     "Αυτόματος υπολογισμός: Κόστος στησίματος ÷ Κυβικά ανά στήσιμο",
                     "Επιβάρυνση αντλίας σε €/m³", yellow=False)
    put_master_price("Τ.Β.7", f"ΤΕΛΙΚΗ ΤΙΜΗ ΣΚΥΡΟΔΕΤΗΣΗΣ ({cls}) ΑΝΑ ΚΥΒΙΚΟ [ΣΤΡΟΓΓΥΛΟΠΟΙΗΜΕΝΗ ΣΕ ΑΚΕΡΑΙΟ €]", "€",
                     "=ROUND({p.conc_buy}+{p.conc_place}+{p.pump_per_m3},0)", "p.conc",
                     "ROUND(Τιμή Αγοράς + 15,00 € + Στήσιμο Αντλίας/m³, 0)",
                     "Στρογγυλοποιημένη τελική τιμή — Τροφοδοτεί Φύλλο 1 (Γρ. 3) & Φύλλο 2 (2Β.3)",
                     yellow=False, green=True, fmt_override=FMT["€0"])

    # --- ΠΙΝΑΚΑΣ Γ: ΜΟΝΟ ΤΟ SLEEVE ΤΗΣ ΣΥΓΚΕΚΡΙΜΕΝΗΣ ΔΟΥΛΕΙΑΣ (ΜΕ × 1,30 OVERHEAD, PROFIT & DAMAGES)
    r3 += 1
    b.put(ws3, f"A{r3}", f"Γ. ΚΟΣΤΟΛΟΓΗΣΗ SLEEVE ΤΗΣ ΣΥΓΚΕΚΡΙΜΕΝΗΣ ΔΟΥΛΕΙΑΣ — ΜΟΝΟ ΓΙΑ {sl_info['type_label'].upper()}", font=ST["sub_font"])
    r3 += 1
    for i, lab in enumerate(("Κωδ.", "Στοιχείο Κόστους Sleeve Έργου", "Μονάδα", "Τιμή", "Προέλευση / Τύπος", "Κανόνας Κοπής, Φύρας, Εργατικών & ×1,30")):
        b.put(ws3, f"{get_column_letter(1 + i)}{r3}", lab, font=ST["hdr_font"], fill=ST["hdr_fill"], align=CENTER)
    ws3.row_dimensions[r3].height = 26
    r3 += 1

    put_master_price("Τ.Γ.1", f"Εμπορικό μήκος σωλήνα που μας φέρνουν ({sl_info['type_label']})", "m",
                     sl_info["pipe_len"], "sl_pipe_len", sl_info["src_len"],
                     "Μήκος εμπορικού σωλήνα από Τιμοκατάλογο Εταιρείας (Φύλλο 5)", yellow=True)
    put_master_price("Τ.Γ.2", f"Τιμή αγοράς εμπορικού σωλήνα ({sl_info['type_label']})", "€",
                     sl_info["pipe_price"], "sl_pipe_price", sl_info["src_price"],
                     "Τιμή αγοράς ολόκληρου σωλήνα από Τιμοκατάλογο Εταιρείας (Φύλλο 5)", yellow=True, fmt_override=FMT["€4"])
    put_master_price("Τ.Γ.3", "Τιμή υλικού σωλήνα ανά μέτρο (Τιμή σωλήνα ÷ Μήκος σωλήνα)", "€/m",
                     "={sl_pipe_price}/{sl_pipe_len}", "sl_price_m",
                     "Αυτόματος υπολογισμός από Τιμοκατάλογο",
                     "Κόστος υλικού σωλήνα ανά τρέχον μέτρο", yellow=False, fmt_override=FMT["€/m"])
    put_master_price("Τ.Γ.4", "Απαιτούμενο μήκος sleeve ανά σίδερο (Starter + Γωνία πάνω)", "m",
                     "={g0.starter}+{g0.top_bend}", "sl_req_len",
                     "Από Τεχνικούς Υπολογισμούς (1.3.8 Starter + 1.3.9 Γωνία πάνω)",
                     "Μήκος sleeve που απαιτείται να κοπεί για κάθε κάθετη ράβδο", yellow=False)
    put_master_price("Τ.Γ.5", "Χρεωστέο μήκος σωλήνα ανά σίδερο (με στρογγυλοποίηση προς τα πάνω λόγω φύρας)", "m",
                     "=IF({sl_req_len}<=0,0,IF({sl_req_len}<={sl_pipe_len},{sl_pipe_len}/INT(ROUND({sl_pipe_len}/{sl_req_len},6)),CEILING({sl_req_len},{sl_pipe_len})))",
                     "sl_billed_len",
                     "Κανόνας φύρας: αν δεν διαιρεί ακριβώς τον σωλήνα, το ρετάλι χάνεται και χρεώνεται!",
                     "Αν L_req=1.00m σε 3m σωλήνα → 3 τεμ./σωλήνα → χρεωστέο 1.00m (μηδενική φύρα)", yellow=False)
    put_master_price("Τ.Γ.6", "Κόστος υλικού σωλήνα ανά σίδερο (Χρεωστέο μήκος × Τιμή ανά μέτρο)", "€",
                     "={sl_billed_len}*{sl_price_m}", "sl_mat_cost",
                     "Χρεωστέο μήκος σωλήνα × Κόστος σωλήνα ανά μέτρο",
                     "Καθαρό κόστος υλικού σωλήνα ανά κάθετο σίδερο", yellow=False, fmt_override=FMT["€4"])
    put_master_price("Τ.Γ.7", "Εργατικά τοποθέτησης ανά σίδερο (Συνεργείο 1,5 ατόμων × 2 λεπτά × 15,60 €/ώρα)", "€",
                     prices["sl_labor"], "sl_labor", pb.source["sl_labor"],
                     "Σταθερό ανά σίδερο ανεξαρτήτως μήκους/πάχους: (2/60) × 1,5 × 15,60 € = 0,7800 €", yellow=True, fmt_override=FMT["€4"])
    put_master_price("Τ.Γ.8", "Κόστος Τέλλας (Μονωτικής Ταινίας) ανά σίδερο (Ρολό 2,19 € ÷ 25 σίδερα)", "€",
                     prices["sl_tape"], "sl_tape", pb.source["sl_tape"],
                     "Σταθερό ανά σίδερο ανεξαρτήτως μήκους/πάχους: 2,19 € / 25 = 0,0876 €", yellow=True, fmt_override=FMT["€4"])
    put_master_price("Τ.Γ.9", "Εσωτερικό Κόστος Sleeve για εμάς ανά σίδερο (Υλικό + Εργατικά + Τέλλα)", "€",
                     "={sl_mat_cost}+{sl_labor}+{sl_tape}", "sl_base_cost",
                     "Άθροισμα: Υλικό σωλήνα + 0,7800 € (Εργατικά) + 0,0876 € (Τέλλα)",
                     "Καθαρό εσωτερικό κόστος πριν το Overhead, Profit & Damages", yellow=False, fmt_override=FMT["€4"])
    put_master_price("Τ.Γ.10", "Συντελεστής Overhead, Profit & Damages (× 1,30)", "×",
                     prices["sl_markup"], "sl_markup", pb.source["sl_markup"],
                     "Πολλαπλασιάζεται πάντα × 1,30 (+30% Overhead, Κέρδος & Ζημιές) στην τελική τιμή", yellow=True, fmt_override=FMT["×"])
    put_master_price("Τ.Γ.11", f"ΤΕΛΙΚΗ ΤΙΜΗ SLEEVE ΕΡΓΟΥ ({sl_info['short_type']}) ΑΝΑ ΣΙΔΕΡΟ [(Υλικό + Εργατικά + Τέλλα) × 1,30]", "€",
                     "={sl_base_cost}*{sl_markup}", "p.sleeve",
                     "Εσωτερικό Κόστος Sleeve × 1,30 (Overhead, Profit & Damages)",
                     "Τροφοδοτεί αυτόματα Φύλλο 1 (Γρ. 5) & Φύλλο 2 (2Β.5)", yellow=False, green=True, fmt_override=FMT["€4"])

    # --- ΠΙΝΑΚΑΣ Δ: ΣΥΝΟΔΕΙΑ ΑΣΤΥΝΟΜΙΑΣ (ΠΑΣΣΑΛΟΙ ΘΕΜΕΛΙΩΣΗΣ / ΝΕΦΚΕΣ) & ΕΛΕΓΧΟΣ ΣΥΝΟΛΩΝ
    r3 += 1
    b.put(ws3, f"A{r3}", "Δ. ΣΥΝΟΔΕΙΑ ΑΣΤΥΝΟΜΙΑΣ ΓΕΩΤΡΥΠΑΝΟΥ (8 €/km) & ΖΩΝΤΑΝΟΣ ΕΛΕΓΧΟΣ ΣΥΝΟΛΩΝ", font=ST["sub_font"])
    r3 += 1
    police_km_init = proj.get("police_km")
    put_master_price("Τ.Δ.1", "Χιλιόμετρα Συνοδείας Αστυνομίας (Πάσσαλοι Θεμελίωσης / Νέφκες — βάλε 0 αν δεν έχει)", "km",
                     float(police_km_init) if police_km_init is not None else "={dist_km}", "police_km",
                     "Συνδεδεμένο με την απόσταση Τ.1 (ή γράψε 0 αν δεν απαιτείται συνοδεία)",
                     "Πολλαπλασιάζεται με τα 8,00 €/km για τη μεταφορά του γεωτρύπανου", yellow=True)
    put_master_price("Τ.Δ.2", "Κόστος Συνοδείας Αστυνομίας ανά χιλιόμετρο (8,00 €/km)", "€/km",
                     prices["police_km"], "p_police_km", pb.source["police_km"],
                     "Από Τιμοκατάλογο Εταιρείας (Φύλλο 6): 8,00 €/km", yellow=True, fmt_override=FMT["€/km"])
    put_master_price("Τ.Δ.3", "Συνολικό Κόστος Συνοδείας Αστυνομίας (Χιλιόμετρα × 8,00 €/km)", "€",
                     "={police_km}*{p_police_km}", "mob_police",
                     "Αυτόματος υπολογισμός: Χιλιόμετρα Συνοδείας × 8,00 €/km",
                     "Προστίθεται αυτόματα στις Μετακινήσεις (Φύλλο 1 Γρ. 1 & Φύλλο 2 2Β.1)", yellow=False, green=True)

    r3 += 1
    ws3.merge_cells(f"A{r3}:C{r3}")
    b.put(ws3, f"A{r3}", "ΓΕΝΙΚΟ ΣΥΝΟΛΟ ΠΡΟΣΦΟΡΑΣ (ΧΩΡΙΣ Φ.Π.Α.):", font=ST["gt_font"], fill=ST["gold"], align=RIGHT, border=ST["gt_border"])
    b.put(ws3, f"D{r3}", "={grand_total}", name="s3.grand_total", fmt=FMT["€"], font=ST["gt_font"], fill=ST["gold"], align=RIGHT, border=ST["gt_border"])
    ws3.merge_cells(f"E{r3}:F{r3}")
    b.put(ws3, f"E{r3}", '="Μέσο Κόστος ανά Πάσσαλο: "&ROUND({grand_total}/{n_piles},2)&" € ("&{n_piles}&" πάσσαλοι)"',
          font=ST["gt_font"], fill=ST["gold"], align=LEFT_WRAP, border=ST["gt_border"])
    ws3.row_dimensions[r3].height = 28

    for c_let, wdt in zip("ABCDEF", (10, 58, 12, 22, 38, 52)):
        ws3.column_dimensions[c_let].width = wdt

    # =========================================================================
    # ΦΥΛΛΟ 2: «Τεχνικοί Υπολογισμοί»
    # =========================================================================
    COL0 = 4
    gcol = [get_column_letter(COL0 + i) for i in range(n)]
    tcol = get_column_letter(COL0 + n)
    ncol = get_column_letter(COL0 + n + 1)
    last = get_column_letter(max(COL0 + n + 1, 7))

    def title(ws, row, text, font=ST["sub_font"]):
        b.put(ws, f"A{row}", text, font=font)

    def header(row, labels):
        for i, lab in enumerate(labels):
            b.put(ws2, f"{get_column_letter(1 + i)}{row}", lab, font=ST["hdr_font"], fill=ST["hdr_fill"], align=CENTER)
        ws2.row_dimensions[row].height = 30

    def scalar(row, code, desc, unit, value, name, note="", red=False, yellow=False, linked=False):
        b.put(ws2, f"A{row}", code, align=CENTER, border=ST["thin"])
        b.put(ws2, f"B{row}", desc, border=ST["thin"])
        b.put(ws2, f"C{row}", unit, align=CENTER, border=ST["thin"])
        c = b.put(ws2, f"D{row}", value, name=name, fmt=FMT.get(unit), align=RIGHT, border=ST["thin"])
        if yellow:
            c.fill = ST["yellow"]
            c.font = ST["bold"]
        elif linked:
            c.fill = ST["linked_fill"]
            c.font = ST["bold"]
        if red:
            c.fill, c.font, c.border = ST["red_fill"], ST["red_font"], ST["red_border"]
        b.put(ws2, f"{ncol}{row}", note, font=ST["note_font"], align=LEFT_WRAP)

    ws2.merge_cells(f"A1:{last}1")
    b.put(ws2, "A1", "ΘΕΜΕΛΙΟΤΕΧΝΙΚΗ ΛΤΔ — ΤΕΧΝΙΚΟΙ ΥΠΟΛΟΓΙΣΜΟΙ ΔΙΑΤΡΗΤΩΝ ΠΑΣΣΑΛΩΝ",
          font=ST["title_font"], fill=ST["title_fill"], align=CENTER)
    ws2.row_dimensions[1].height = 28
    b.put(ws2, "A2", f"Έργο: {proj.get('name', '')} | Πελάτης: {proj.get('client', '')} | "
                     f"Μελέτη: {proj.get('designer', '')} | Μηχανή {ENGINE_VERSION} (Τιμές από Φύλλο 3: «{S3}»)",
          font=ST["meta_font"])

    r = 4
    title(ws2, r, "1. ΔΕΔΟΜΕΝΑ ΣΧΕΔΙΟΥ & ΥΨΟΜΕΤΡΑ")
    r += 1
    header(r, ["Κωδ.", "Παράμετρος", "Μονάδα", "Τιμή"] + [""] * n + ["Παρατηρήσεις / Προέλευση"])
    r += 1
    title(ws2, r, "Τομέας 1 — Μετακινήσεις & Συνοδεία Αστυνομίας (Τιμές από το Φύλλο 3: «Τιμές & Έλεγχος»)")
    r += 1
    scalar(r, "1.1.1", "Χιλιόμετρα εκτός Λεμεσού (0 km αν το έργο είναι μέσα στη Λεμεσό)", "km",
           "={dist_km}", "s2.dist_km", "Από Φύλλο 3 («Τιμές & Έλεγχος» Τ.1)", linked=True); r += 1
    scalar(r, "1.1.2", "Βήμα χιλιομετρικής χρέωσης εκτός Λεμεσού (ανά 10 km)", "km",
           consts["zone_km"], "zone_km", pb.source["zone_km"]); r += 1
    scalar(r, "1.1.3", "Δεκάδες χιλιομέτρων εκτός Λεμεσού (ανά 10 km)", "ζώνες",
           "=IF({dist_km}>0,ROUNDUP(ROUND({dist_km}/{zone_km},9),0),0)", "extra_zones",
           "ROUNDUP(km εκτός Λεμεσού / 10, 0)"); r += 1
    scalar(r, "1.1.4", "Μετακίνηση — βασική τιμή εντός Λεμεσού", "€", "={p_mob_base}", "s2.p_mob_base",
           "Από Φύλλο 3 («Τιμές & Έλεγχος» Τ.2 = 1.000,00 €)", linked=True); r += 1
    scalar(r, "1.1.5", "Μετακίνηση — επιβάρυνση ανά 10 km εκτός Λεμεσού", "€", "={p_mob_extra}", "s2.p_mob_extra",
           "Από Φύλλο 3 («Τιμές & Έλεγχος» Τ.3 = 200,00 €/10km)", linked=True); r += 1
    scalar(r, "1.1.6", "Κόστος Συνοδείας Αστυνομίας (Χιλιόμετρα × 8,00 €/km)", "€", "={mob_police}", "s2.mob_police",
           "Από Φύλλο 3 («Τιμές & Έλεγχος» Τ.Δ.3 = police_km × 8,00 €/km)", linked=True); r += 1
    scalar(r, "1.1.7", "Συνολικό εφάπαξ ποσό μετακίνησης & συνοδείας Αστυνομίας", "€",
           "={p_mob_base}+{extra_zones}*{p_mob_extra}+{mob_police}", "mob_total",
           "Βασική Λεμεσού (1.000€) + Δεκάδες km × 200€ + Συνοδεία Αστυνομίας"); r += 1
    scalar(r, "1.1.8", "Συνολικό πλήθος πασσάλων έργου", "τεμ.",
           "=" + "+".join(f"{{g{i}.count}}" for i in range(n)), "n_piles", "Άθροισμα ομάδων"); r += 1
    scalar(r, "1.1.9", "Επιμερισμένη μετακίνηση ανά πάσσαλο", "€",
           "={mob_total}/{n_piles}", "mob_per_pile", "Συνολική μετακίνηση ÷ Πάσσαλοι"); r += 1

    title(ws2, r, "Τομέας 2 — Υψόμετρα & Σκυρόδεμα (WPL κοινό για όλους, Cut-Off ανά πάσσαλο στον Τομέα 3)"); r += 1
    wpl_note = ("ΚΟΚΚΙΝΟΣ ΣΥΝΑΓΕΡΜΟΣ: δεν δόθηκε — προσωρινά 0.00" if el.get("wpl") is None
                else el.get("wpl_note", "Από σχέδιο / χρήστη"))
    scalar(r, "1.2.1", "Στάθμη πλατφόρμας γεωτρύπανου (WPL) — κοινή για όλους τους πασσάλους", "m",
           float(el["wpl"]) if el.get("wpl") is not None else 0.0, "wpl", wpl_note, red=el.get("wpl") is None); r += 1
    scalar(r, "1.2.2", "Κατηγορία σκυροδέματος", "-", cls, "conc_class",
           spec.get("concrete_note", "")); r += 1
    scalar(r, "1.2.3", "Επικάλυψη οπλισμού (Cover)", "mm", float(spec.get("cover_mm", 75)), "cover_mm",
           spec.get("cover_note", "")); r += 1

    title(ws2, r, "Τομέας 3 — Ομάδες πασσάλων & κλωβός οπλισμού (μία στήλη ανά ομάδα)"); r += 1
    header(r, ["Κωδ.", "Παράμετρος", "Μονάδα"] + [f"Ομάδα {i + 1}" for i in range(n)] + ["", "Παρατηρήσεις"])
    r += 1
    b.put(ws2, f"A{r}", "1.3.0", align=CENTER, border=ST["thin"])
    b.put(ws2, f"B{r}", "Πάσσαλοι", border=ST["thin"])
    for i, g in enumerate(groups):
        b.put(ws2, f"{gcol[i]}{r}", f"{g['label']}\n{g['ids']}", align=CENTER, border=ST["thin"], font=ST["bold"])
    ws2.row_dimensions[r].height = 30
    r += 1
    for k, (key, desc, unit, _req, _d) in enumerate(GROUP_INPUTS, 1):
        b.put(ws2, f"A{r}", f"1.3.{k}", align=CENTER, border=ST["thin"])
        b.put(ws2, f"B{r}", desc, border=ST["thin"])
        b.put(ws2, f"C{r}", unit, align=CENTER, border=ST["thin"])
        for i, g in enumerate(groups):
            val = g[key]
            red = False
            if key == "cutoff" and val is None:
                val, red = "ΛΕΙΠΕΙ", True
            elif key == "starter" and val is None:
                val = "={G.raft}-{c_starter_off}".replace("{G.", f"{{g{i}.")
            elif key == "top_bend" and val is None:
                val = "={c_top_bend}"
            elif key == "lap_len" and val is None:
                val = "={c_lap}*{G.y_vert}/1000".replace("{G.", f"{{g{i}.")
            elif key in ("raft", "head_links") and val is None:
                val = "—"
            c = b.put(ws2, f"{gcol[i]}{r}", val, name=f"g{i}.{key}", fmt=FMT.get(unit), align=RIGHT, border=ST["thin"])
            if red:
                c.fill, c.font, c.border = ST["red_fill"], ST["red_font"], ST["red_border"]
        if key == "count":
            b.put(ws2, f"{tcol}{r}", "=" + "+".join(f"{{g{i}.count}}" for i in range(n)), name="t.count",
                  fmt=FMT["τεμ."], align=RIGHT, border=ST["thin"], font=ST["bold"])
        note = (spec.get("input_notes") or {}).get(key, "")
        b.put(ws2, f"{ncol}{r}", note, font=ST["note_font"], align=LEFT_WRAP)
        r += 1

    # --- Ενότητα 2: Σμαραγδί πίνακας
    r += 1
    title(ws2, r, "2. ΤΕΛΙΚΕΣ ΠΟΣΟΤΗΤΕΣ ΓΙΑ ΤΗΝ ΟΙΚΟΝΟΜΙΚΗ ΠΡΟΣΦΟΡΑ", font=ST["em_title"]); r += 1
    em_hdr = ["Κωδ.", "Ποσότητα", "Μονάδα", "Σύνολο Έργου", "Πλήθος Πασσάλων", "Μέση ανά Πάσσαλο", "Τροφοδοτεί"]
    for i, lab in enumerate(em_hdr):
        b.put(ws2, f"{get_column_letter(1 + i)}{r}", lab, font=ST["hdr_font"], fill=ST["hdr_fill"], align=CENTER)
    ws2.row_dimensions[r].height = 30
    r += 1
    em_rows = []
    for d in diameters:
        idx = [i for i, g in enumerate(groups) if g["D_mm"] == d]
        em_rows.append((f"drill_{d}", f"Τρέχοντα μέτρα διάτρησης Ø{d}mm", "m",
                        "=" + "+".join(f"{{g{i}.count}}*{{g{i}.L_drill}}" for i in idx),
                        "=" + "+".join(f"{{g{i}.count}}" for i in idx)))
    em_rows += [
        ("conc", '="Σκυρόδεμα "&{conc_class}&" με απώλεια"', "m³", "={t.V_conc}", "={n_piles}"),
        ("steel", "Χάλυβας οπλισμού B500C (κάθετα σίδερα + σούστα + στεφάνια αναμονής)", "kg", "={t.W_total}", "={n_piles}"),
        ("sleeve", f'="Sleeves προστασίας αναμονών starter ({sl_info["short_type"]} για Φ"&{{g0.y_vert}}&")"', "τεμ.", "={t.sleeves}", "={n_piles}"),
        ("breaking", "Κεφαλοκοπή πασσάλων (σπάσιμο κεφαλής)", "τεμ.", "={t.breaking}", "={n_piles}"),
    ]
    for k, (key, desc, unit, f_tot, f_cnt) in enumerate(em_rows, 1):
        vals = [f"2.{k}", desc, unit]
        for i, v in enumerate(vals):
            b.put(ws2, f"{get_column_letter(1 + i)}{r}", v)
        b.put(ws2, f"D{r}", f_tot, name=f"q.{key}", fmt=FMT.get(unit))
        b.put(ws2, f"E{r}", f_cnt, name=f"q.{key}.n", fmt=FMT["τεμ."])
        b.put(ws2, f"F{r}", f"={{q.{key}}}/{{q.{key}.n}}", fmt=FMT.get(unit))
        b.put(ws2, f"G{r}", "Φύλλο 1")
        for c in range(1, 8):
            cell = ws2.cell(r, c)
            cell.fill, cell.font, cell.border = ST["em_fill"], ST["em_font"], ST["em_border"]
        r += 1

    # --- Ενότητα 2Β: Συγκεντρωτικός Πίνακας Πράξεων, Ποσοτήτων & Τιμών
    r += 1
    title(ws2, r, f"2Β. ΣΥΓΚΕΝΤΡΩΤΙΚΟΣ ΠΙΝΑΚΑΣ ΠΡΑΞΕΩΝ, ΠΟΣΟΤΗΤΩΝ & ΤΙΜΩΝ (ΤΙΜΕΣ ΑΠΟ ΦΥΛΛΟ 3: «{S3}»)",
          font=ST["sub_font"])
    r += 1
    b2_hdr = [
        "Κωδ.", "Πράξη / Εργασία (Εκσκαφή, Σκυρόδεμα, Οπλισμός κ.λπ.)", "Μονάδα",
        "Συνολική Ποσότητα", f"Τιμή Μονάδας (€)\n[Από «{S3}»]", "Συνολικό Κόστος (€)",
        "Ανθρώπινη Επεξήγηση Πράξεων (Για Έλεγχο από Μηχανικό / Μάστρο — 100% Δυναμική)"
    ]
    for i, lab in enumerate(b2_hdr):
        b.put(ws2, f"{get_column_letter(1 + i)}{r}", lab, font=ST["hdr_font"], fill=ST["hdr_fill"], align=CENTER)
    if last != "G":
        ws2.merge_cells(f"G{r}:{last}{r}")
    ws2.row_dimensions[r].height = 30
    r += 1

    b2_rows = build_consolidated_rows(spec, groups, diameters, consts, prices, cls, ref, sl_info)
    b2_main_keys = [row["key"] for row in b2_rows if row["is_main"]]
    for row in b2_rows:
        rk = row["key"]
        is_main = row["is_main"]
        row_font = ST["bold"] if is_main else ST["sub_row_font"]
        row_fill = None if is_main else ST["sub_row_fill"]
        p_fmt = FMT["€4"] if rk == "sleeve" else (FMT["€0"] if rk == "conc" else FMT["€"])

        b.put(ws2, f"A{r}", row["code"], align=CENTER, border=ST["thin"], font=row_font, fill=row_fill)
        b.put(ws2, f"B{r}", row["desc"], align=LEFT_WRAP, border=ST["thin"], font=row_font, fill=row_fill)
        b.put(ws2, f"C{r}", row["unit"], align=CENTER, border=ST["thin"], font=row_font, fill=row_fill)
        b.put(ws2, f"D{r}", row["f_qty"], name=f"b2.{rk}.qty", fmt=FMT.get(row["unit"], "#,##0.00"),
              align=RIGHT, border=ST["thin"], font=row_font, fill=row_fill)
        b.put(ws2, f"E{r}", row["f_price"], name=f"b2.{rk}.price", fmt=p_fmt,
              align=RIGHT, border=ST["thin"], font=row_font, fill=ST["linked_fill"] if is_main else row_fill)
        b.put(ws2, f"F{r}", f"={{b2.{rk}.qty}}*{{b2.{rk}.price}}", name=f"b2.{rk}.total", fmt=FMT["€"],
              align=RIGHT, border=ST["thin"], font=row_font, fill=row_fill)
        if last != "G":
            ws2.merge_cells(f"G{r}:{last}{r}")
        b.put(ws2, f"G{r}", row["math"], name=f"b2.{rk}.note", align=LEFT_WRAP, border=ST["thin"],
              font=ST["note_font"], fill=row_fill)
        ws2.row_dimensions[r].height = 34
        r += 1

    # Γενικό Σύνολο Πίνακα 2Β
    ws2.merge_cells(f"A{r}:E{r}")
    b.put(ws2, f"A{r}", "ΓΕΝΙΚΟ ΣΥΝΟΛΟ ΕΡΓΟΥ (ΑΘΡΟΙΣΜΑ ΚΥΡΙΩΝ ΠΡΑΞΕΩΝ 2Β.1 ΕΩΣ 2Β.6 — ΧΩΡΙΣ Φ.Π.Α.)",
          font=ST["gt_font"], align=RIGHT)
    b.put(ws2, f"F{r}", "=" + "+".join(f"{{b2.{k}.total}}" for k in b2_main_keys), name="b2.grand_total",
          fmt=FMT["€"], font=ST["gt_font"], align=RIGHT)
    if last != "G":
        ws2.merge_cells(f"G{r}:{last}{r}")
    b.put(ws2, f"G{r}",
          '="Σύνολο Προσφοράς = "&ROUND({b2.grand_total},2)&" € | Μέσο Κόστος ανά Πάσσαλο = "'
          '&ROUND({b2.grand_total}/{n_piles},2)&" €/πάσσαλο ("&{n_piles}&" πάσσαλοι)"',
          name="b2.grand_total.note", font=ST["gt_font"], align=LEFT_WRAP)
    for cc in range(1, max(COL0 + n + 1, 7) + 1):
        ws2.cell(r, cc).fill = ST["gold"]
        ws2.cell(r, cc).border = ST["gt_border"]
    ws2.row_dimensions[r].height = 26
    r += 1

    # --- Ενότητα 3: Σταθερές (αριθμοί στη στήλη D)
    r += 1
    title(ws2, r, "3. ΣΤΑΘΕΡΕΣ ΤΕΧΝΙΚΕΣ ΠΑΡΑΔΟΧΕΣ (αριθμοί στη Στήλη D)"); r += 1
    header(r, ["Κωδ.", "Παραδοχή", "Μονάδα", "Τιμή"] + [""] * n + ["Προέλευση"]); r += 1
    scalar(r, "3.1", "Συντελεστής απώλειας σκυροδέματος", "%", consts["c_loss"], "c_loss", pb.source["c_loss"]); r += 1
    scalar(r, "3.2", "Πυκνότητα χάλυβα ρ (kg/m³)", "-", consts["c_density"], "c_density", pb.source["c_density"])
    ws2[f"D{r}"].number_format = "#,##0"; r += 1
    scalar(r, "3.3", "Μέγιστο εμπορικό μήκος βέργας", "m", consts["c_barlen"], "c_barlen", pb.source["c_barlen"]); r += 1
    scalar(r, "3.4", "Συντελεστής μάτισης (× Φ, αν δεν δίνεται στο σχέδιο)", "-", consts["c_lap"], "c_lap", pb.source["c_lap"])
    ws2[f"D{r}"].number_format = "0"; r += 1
    scalar(r, "3.5", "Μέγιστο ύψος υπερκάλυψης σκυροδέματος", "m", consts["c_overc"], "c_overc", pb.source["c_overc"]); r += 1
    scalar(r, "3.6", "Μαθηματική σταθερά π", "-", "=PI()", "c_pi", "PI()"); r += 1
    scalar(r, "3.7", "Starter = βάθος πεδιλόπλακας − (αν δεν δίνεται starter)", "m", RULE_STARTER_OFFSET,
           "c_starter_off", "Κανόνας εταιρείας RULES_v5 §11.1"); r += 1
    scalar(r, "3.8", "Γωνία πάνω (αν δεν δίνεται στο σχέδιο)", "m", RULE_TOP_BEND, "c_top_bend",
           "Κανόνας εταιρείας RULES_v5 §11.2"); r += 1

    # --- Ενότητα 4: Αναλυτικά μαθηματικά (γραμμές = βήματα, στήλες = ομάδες)
    r += 1
    title(ws2, r, "4. ΑΝΑΛΥΤΙΚΑ ΜΑΘΗΜΑΤΙΚΑ ΒΗΜΑ-ΒΗΜΑ (τιμές ανά πάσσαλο, σύνολο = Σ πλήθος × τιμή)"); r += 1
    header(r, ["Βήμα", "Υπολογισμός", "Μονάδα"] + [f"Ομάδα {i + 1}\n{groups[i]['label']}" for i in range(n)]
           + ["Σύνολο Έργου", "Τρόπος Υπολογισμού (Ανθρώπινα)"])
    r += 1
    for k, (key, desc, unit, tpl, additive, text) in enumerate(MATH_STEPS, 1):
        b.put(ws2, f"A{r}", str(k), align=CENTER, border=ST["thin"])
        b.put(ws2, f"B{r}", desc, border=ST["thin"])
        b.put(ws2, f"C{r}", unit, align=CENTER, border=ST["thin"])
        for i in range(n):
            tpl_i = tpl
            dh_missing = key == "dh" and (el.get("wpl") is None or groups[i]["cutoff"] is None)
            if key == "n_head" and groups[i]["head_links"] is not None:
                tpl_i = "={G.head_links}"
            if dh_missing:
                tpl_i = 0.0
            val_i = tpl_i.replace("{G.", f"{{g{i}.") if isinstance(tpl_i, str) else tpl_i
            c = b.put(ws2, f"{gcol[i]}{r}", val_i, name=f"g{i}.{key}",
                      fmt=FMT.get(unit), align=RIGHT, border=ST["thin"])
            if dh_missing:
                c.fill, c.font, c.border = ST["red_fill"], ST["red_font"], ST["red_border"]
        if additive:
            b.put(ws2, f"{tcol}{r}", "=" + "+".join(f"{{g{i}.count}}*{{g{i}.{key}}}" for i in range(n)),
                  name=f"t.{key}", fmt=FMT.get(unit), align=RIGHT, border=ST["thin"], font=ST["bold"])
        else:
            b.put(ws2, f"{tcol}{r}", "—", align=CENTER, border=ST["thin"])
        b.put(ws2, f"{ncol}{r}", text, font=ST["note_font"], align=LEFT_WRAP)
        r += 1

    # --- Ενότητα 5: Πηγές δεδομένων (Drawing Audit)
    r += 1
    title(ws2, r, "5. ΠΗΓΕΣ ΔΕΔΟΜΕΝΩΝ — ΑΠΟ ΠΟΥ ΠΡΟΗΛΘΕ ΚΑΘΕ ΣΤΟΙΧΕΙΟ (DRAWING AUDIT)"); r += 1
    src_hdr_last = get_column_letter(max(COL0 + n + 1, 7))
    for col, lab in (("A", "Α/Α"), ("B", "Στοιχείο"), ("C", "Κατάσταση"), ("D", "Τιμή"), ("E", "Πηγή / Σημείο στο σχέδιο / Αιτιολόγηση")):
        b.put(ws2, f"{col}{r}", lab, font=ST["hdr_font"], fill=ST["hdr_fill"], align=CENTER)
    ws2.merge_cells(f"E{r}:{src_hdr_last}{r}")
    r += 1
    STATUS_FILL = {"ΣΧΕΔΙΟ": PatternFill("solid", start_color="DCFCE7"),
                   "ΧΡΗΣΤΗΣ": PatternFill("solid", start_color="DBEAFE"),
                   "ΤΙΜΟΚΑΤΑΛΟΓΟΣ": PatternFill("solid", start_color="F3F4F6"),
                   "ΜΕΤΡΗΣΗ": PatternFill("solid", start_color="FEF9C3"),
                   "ΚΑΝΟΝΙΣΜΟΣ": PatternFill("solid", start_color="EDE9FE"),
                   "ΚΑΝΟΝΑΣ": PatternFill("solid", start_color="FFEDD5"),
                   "ΠΑΡΑΔΟΧΗ": ST["red_fill"]}
    sources = spec.get("sources") or [{"item": "—", "status": "ΠΑΡΑΔΟΧΗ", "value": "—",
                                       "source": "ΔΕΝ ΔΗΛΩΘΗΚΑΝ ΠΗΓΕΣ ΣΤΟ SPEC — η προσφορά δεν πρέπει να σταλεί."}]
    for k, s in enumerate(sources, 1):
        st = s.get("status", "")
        b.put(ws2, f"A{r}", str(k), align=CENTER, border=ST["thin"])
        b.put(ws2, f"B{r}", s.get("item", ""), align=LEFT_WRAP, border=ST["thin"])
        b.put(ws2, f"C{r}", st, align=CENTER, border=ST["thin"], fill=STATUS_FILL.get(st), font=ST["bold"])
        b.put(ws2, f"D{r}", str(s.get("value", "")), align=CENTER, border=ST["thin"])
        ws2.merge_cells(f"E{r}:{src_hdr_last}{r}")
        b.put(ws2, f"E{r}", s.get("source", ""), align=LEFT_WRAP, font=ST["note_font"])
        txt_len = len(s.get("source", ""))
        val_len = len(str(s.get("value", "")))
        item_len = len(s.get("item", ""))
        ws2.row_dimensions[r].height = 15 * max(1, math.ceil(txt_len / 140), math.ceil(val_len / 13),
                                                math.ceil(item_len / 45))
        r += 1
    b.put(ws2, f"A{r}", "Κατάσταση: ΣΧΕΔΙΟ = γράφεται ρητά στο σχέδιο | ΜΕΤΡΗΣΗ = μετρήθηκε γραφικά από το σχέδιο | "
                        "ΚΑΝΟΝΙΣΜΟΣ = Ευρωκώδικας | ΚΑΝΟΝΑΣ = κανόνας εταιρείας (RULES) | ΤΙΜΟΚΑΤΑΛΟΓΟΣ | ΧΡΗΣΤΗΣ = οδηγία χρήστη | "
                        "ΠΑΡΑΔΟΧΗ = δεν δίνεται, απαιτεί επιβεβαίωση", font=ST["meta_font"])
    r += 1

    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 54
    ws2.column_dimensions["C"].width = 10
    for c in gcol:
        ws2.column_dimensions[c].width = 16
    ws2.column_dimensions[tcol].width = 18
    ws2.column_dimensions[ncol].width = 74
    ws2.column_dimensions["G"].width = max(ws2.column_dimensions["G"].width or 0, 74)
    ws2.freeze_panes = "D4"

    # =========================================================================
    # ΦΥΛΛΟ 1: «Οικονομική Προσφορά»
    # =========================================================================
    ws1.merge_cells("A1:G1")
    b.put(ws1, "A1", "ΘΕΜΕΛΙΟΤΕΧΝΙΚΗ ΛΤΔ — ΟΙΚΟΝΟΜΙΚΗ ΠΡΟΣΦΟΡΑ ΠΑΣΣΑΛΩΝ",
          font=ST["title_font"], fill=ST["title_fill"], align=CENTER)
    ws1.row_dimensions[1].height = 28
    b.put(ws1, "A2", f"Έργο: {proj.get('name', '')} | Πελάτης: {proj.get('client', '')} | "
                     f"Τοποθεσία: {proj.get('location', '')} | Τιμές από Φύλλο 3: «{S3}»", font=ST["meta_font"])

    if red_alert:
        ws1.merge_cells("A4:G5")
        b.put(ws1, "A4", "🛑 ΠΡΟΣΟΧΗ — RED ALERT: Δεν έχουν δοθεί η στάθμη πλατφόρμας (WPL, κοινή για όλους) ή/και οι "
                         "στάθμες αποκοπής (Cut-Off) ανά πάσσαλο. Όπου λείπουν, η προσφορά υπολογίζει Δh = 0.00 m "
                         "(χωρίς κενή διάτρηση, sleeves και κεφαλοκοπή). Απαιτείται συμπλήρωση πριν την οριστικοποίηση.",
              font=ST["red_font"], fill=ST["red_fill"], align=CENTER)
        for rr in (4, 5):
            for cc in range(1, 8):
                ws1.cell(rr, cc).border = ST["red_border"]
        ws1.row_dimensions[4].height = 24
        ws1.row_dimensions[5].height = 24
    for extra_i, banner in enumerate(spec.get("extra_banners", [])):
        rr = 6 + extra_i
        ws1.merge_cells(f"A{rr}:G{rr}")
        b.put(ws1, f"A{rr}", banner, font=ST["red_font"], fill=ST["red_fill"], align=CENTER)
        ws1.row_dimensions[rr].height = 30
    r = 7 + len(spec.get("extra_banners", []))
    hdr1 = ["Α/Α", "Περιγραφή Εργασίας / Υλικού", "Ποσότητα", "Μονάδα",
            f"Τιμή Μονάδας (€)\n[Από «{S3}»]",
            "Σύνολο (€)", "Μέσο Κόστος ανά Πάσσαλο (€)"]
    for i, lab in enumerate(hdr1):
        b.put(ws1, f"{get_column_letter(1 + i)}{r}", lab, font=ST["hdr_font"], fill=ST["hdr_fill"], align=CENTER)
    ws1.row_dimensions[r].height = 30
    r += 1

    lines = [("mob", "1", "Μετακινήσεις γεωτρύπανου & Συνοδεία Αστυνομίας (επιμερισμός ανά πάσσαλο)",
              "={n_piles}", "πάσσαλο", "={mob_per_pile}")]
    for j, d in enumerate(diameters):
        idx = [i for i, g in enumerate(groups) if g["D_mm"] == d]
        cnt = sum(groups[i]["count"] for i in idx)
        sub = f"2{chr(ord('a') + j)}" if len(diameters) > 1 else "2"
        lines.append((f"drill_{d}", sub, f"Τρύπημα / διάτρηση πασσάλων Ø{d}mm ({cnt} τεμ.)",
                      f"={{q.drill_{d}}}", "m", f"={{p.drill_{d}}}"))
    lines += [
        ("conc", "3", '="Σκυροδέτηση πασσάλων "&{conc_class}&" (με απώλεια +"&ROUND({c_loss}*100,0)&"%)"',
         "={q.conc}", "m³", "={p.conc}"),
        ("steel", "4",
         '="Χάλυβας οπλισμού B500C (κάθετα σίδερα "&{g0.n_vert}&"Φ"&{g0.y_vert}&", σούστα Φ"&{g0.y_spiral}'
         '&"/"&ROUND({g0.s_mid}/10,0)&" & στεφάνια αναμονής)"',
         "={q.steel}", "kg", "={p.steel}"),
        ("sleeve", "5",
         f'="Sleeves προστασίας αναμονών starter ({sl_info["short_type"]} για κάθετα σίδερα Φ"&{{g0.y_vert}}&")"',
         "={q.sleeve}", "τεμ.", "={p.sleeve}"),
        ("breaking", "6", "Κεφαλοκοπή πασσάλων / Pile Breaking", "={q.breaking}", "τεμ.", "={p.breaking}"),
    ]
    for key, aa, desc, f_qty, unit, f_price in lines:
        p_fmt = FMT["€4"] if key == "sleeve" else (FMT["€0"] if key == "conc" else FMT["€"])
        b.put(ws1, f"A{r}", aa, align=CENTER, border=ST["thin"])
        b.put(ws1, f"B{r}", desc, name=f"line.{key}.desc", align=LEFT_WRAP, border=ST["thin"])
        b.put(ws1, f"C{r}", f_qty, name=f"line.{key}.qty", fmt=FMT.get(unit, "#,##0.00"), align=RIGHT, border=ST["thin"])
        b.put(ws1, f"D{r}", unit, align=CENTER, border=ST["thin"])
        b.put(ws1, f"E{r}", f_price, name=f"line.{key}.price", fmt=p_fmt, align=RIGHT,
              border=ST["thin"], fill=ST["linked_fill"], font=ST["bold"])
        b.put(ws1, f"F{r}", f"={{line.{key}.qty}}*{{line.{key}.price}}", name=f"line.{key}.total",
              fmt=FMT["€"], align=RIGHT, border=ST["thin"])
        b.put(ws1, f"G{r}", f"={{line.{key}.total}}/{{n_piles}}", fmt=FMT["€"], align=RIGHT, border=ST["thin"])
        ws1.row_dimensions[r].height = 30
        r += 1
    ws1.merge_cells(f"A{r}:E{r}")
    b.put(ws1, f"A{r}", "ΓΕΝΙΚΟ ΣΥΝΟΛΟ ΠΡΟΣΦΟΡΑΣ (ΧΩΡΙΣ Φ.Π.Α.)", font=ST["gt_font"], align=RIGHT)
    b.put(ws1, f"F{r}", "=" + "+".join(f"{{line.{k[0]}.total}}" for k in lines), name="grand_total",
          fmt=FMT["€"], font=ST["gt_font"], align=RIGHT)
    b.put(ws1, f"G{r}", "={grand_total}/{n_piles}", fmt=FMT["€"], font=ST["gt_font"], align=RIGHT)
    for cc in range(1, 8):
        ws1.cell(r, cc).fill = ST["gold"]
        ws1.cell(r, cc).border = ST["gt_border"]
    ws1.row_dimensions[r].height = 26
    r += 2

    # --- ΔΥΝΑΜΙΚΟΙ ΟΡΟΙ, ΠΑΡΑΔΟΧΕΣ & ΠΑΡΑΤΗΡΗΣΕΙΣ ΣΤΟ ΦΥΛΛΟ 1 (Ενημερώνονται αυτόματα από το Φύλλο 3!)
    b.put(ws1, f"A{r}", f"ΟΡΟΙ, ΠΑΡΑΔΟΧΕΣ & ΠΑΡΑΤΗΡΗΣΕΙΣ (100% ΔΥΝΑΜΙΚΕΣ ΣΗΜΕΙΩΣΕΙΣ — ΕΝΗΜΕΡΩΝΟΝΤΑΙ ΑΥΤΟΜΑΤΑ ΑΠΟ ΤΟ ΦΥΛΛΟ «{S3}»):",
          font=Font(name="Calibri", size=10, bold=True, color="1E40AF"))
    r += 1

    d_main = diameters[0]
    dyn_notes = [
        (
            "s1.note.1",
            '="1. Στις τιμές ΔΕΝ περιλαμβάνεται ο Φ.Π.Α. — Γενικό Σύνολο Προσφοράς: "&ROUND({grand_total},2)'
            '&" € (Μέσο κόστος ανά πάσσαλο: "&ROUND({grand_total}/{n_piles},2)&" € για "&{n_piles}&" πασσάλους)."'
        ),
        (
            "s1.note.2",
            '="2. Μετακινήσεις γεωτρύπανου: Βασική Λεμεσού ("&ROUND({p_mob_base},2)&" €) + Εκτός Λεμεσού ανά "'
            '&ROUND({zone_km},0)&"km ("&{extra_zones}&" × "&ROUND({p_mob_extra},2)&" € = "'
            '&ROUND({extra_zones}*{p_mob_extra},2)&" €) + Συνοδεία Αστυνομίας ("&ROUND({police_km},1)&" km × "'
            '&ROUND({p_police_km},2)&" €/km = "&ROUND({mob_police},2)&" €) = "&ROUND({mob_total},2)'
            '&" € ("&ROUND({line.mob.price},2)&" €/πάσσαλο)."'
        ),
        (
            "s1.note.3",
            f'="3. Εκσκαφή / Διάτρηση (από στάθμη πλατφόρμας WPL "&ROUND({{wpl}},2)&"m): Ύψος σκυροδετημένου πασσάλου = "'
            f'&ROUND({{g0.L_pile}},3)&"m, Συνολικά μέτρα διάτρησης = "&ROUND({{q.drill_{d_main}}},2)&"m × Τιμή μέτρου ("'
            f'&ROUND({{line.drill_{d_main}.price}},2)&" €) = "&ROUND({{line.drill_{d_main}.total}},2)&" €."'
        ),
        (
            "s1.note.4",
            '="4. Σκυρόδεμα "&{conc_class}&": ROUND(Αγορά "&ROUND({p.conc_buy},2)&" €/m³ + Τοποθέτηση "'
            '&ROUND({p.conc_place},2)&" €/m³ + Αντλία "&ROUND({p.pump_per_m3},2)&" €/m³, 0) = "'
            '&ROUND({line.conc.price},0)&" €/m³ × "&ROUND({q.conc},2)&" m³ = "&ROUND({line.conc.total},2)&" €."'
        ),
        (
            "s1.note.5",
            '="5. Οπλισμός κλωβού ("&{g0.n_vert}&"Φ"&{g0.y_vert}&" + σούστα Φ"&{g0.y_spiral}&"/"&ROUND({g0.s_mid}/10,0)'
            '&"): Αναμονή starter = "&ROUND({g0.starter},2)&"m, Γωνία πάνω = "&ROUND({g0.top_bend},2)'
            '&"m, Γωνία κάτω = "&ROUND({g0.bottom_bend},2)&"m, Ύψος μάτισης = "&ROUND({g0.lap_len},2)'
            '&"m ("&{g0.n_laps}&" μάτιση), Σπείρες σούστας = "&{g0.n_turns}&" (με +1) + "&{g0.n_head}'
            '&" στεφάνια αναμονής (χωρίς +1) → "&ROUND({g0.W_total},2)&" kg/πασσ., Σύνολο = "'
            '&ROUND({q.steel},1)&" kg × Τιμή κιλού ("&ROUND({line.steel.price},2)&" €) = "&ROUND({line.steel.total},2)&" €."'
        ),
        (
            "s1.note.6",
            f'="6. Sleeves αναμονών ({sl_info["short_type"]} για Φ"&{{g0.y_vert}}&"): [Χρεωστέο μήκος "'
            f'&ROUND({{sl_billed_len}},2)&"m × "&ROUND({{sl_price_m}},4)&" €/m + Εργατικά ("&ROUND({{sl_labor}},4)'
            f'&" €) + Τέλλα ("&ROUND({{sl_tape}},4)&" €) = "&ROUND({{sl_base_cost}},4)&" €] × "'
            f'&ROUND({{sl_markup}},2)&" (Overhead/Profit/Damages) = "&ROUND({{line.sleeve.price}},4)'
            f'&" €/τεμ. × "&{{q.sleeve}}&" τεμ. = "&ROUND({{line.sleeve.total}},2)&" € | Κεφαλοκοπή: "'
            f'&{{q.breaking}}&" τεμ. × "&ROUND({{line.breaking.price}},2)&" € = "&ROUND({{line.breaking.total}},2)&" €."'
        ),
        (
            "s1.note.7",
            f"7. Πληρωμή βάσει επιμετρήσεων πραγματικά εκτελεσθέντων μέτρων διάτρησης, κυβικών σκυροδέματος και κιλών οπλισμού. "
            f"Όλες οι τιμές μονάδας ορίζονται στο Φύλλο 3 («{S3}») και ενημερώνουν αυτόματα όλα τα φύλλα και τις σημειώσεις."
        ),
    ]
    for n_name, f_txt in dyn_notes:
        ws1.merge_cells(f"A{r}:G{r}")
        b.put(ws1, f"A{r}", f_txt, name=n_name, font=ST["note_font"], align=LEFT_WRAP)
        ws1.row_dimensions[r].height = 24
        r += 1

    for c, wdt in zip("ABCDEFG", (7, 58, 14, 10, 18, 18, 20)):
        ws1.column_dimensions[c].width = wdt

    b.finalize()
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    b.wb.save(out_path)
    return b.layout(), ref, [l[0] for l in lines], pb.source


# =============================================================================
# 7. VERIFICATION GATE (Απομονωμένο DispatchEx — ΔΕΝ αγγίζει ποτέ το Excel του χρήστη)
# =============================================================================
XL_ERRORS = {-2146826288: "#NULL!", -2146826281: "#DIV/0!", -2146826273: "#VALUE!", -2146826265: "#REF!",
             -2146826259: "#NAME?", -2146826252: "#NUM!", -2146826246: "#N/A"}


def verify(path, layout, ref, line_keys):
    try:
        import tempfile
        import win32com
        gp = os.path.join(tempfile.gettempdir(), "pile_engine_genpy")
        os.makedirs(gp, exist_ok=True)
        win32com.__gen_path__ = gp
        import win32com.client
    except Exception as e:  # pragma: no cover
        return {"status": "UNVERIFIED", "reason": f"Δεν υπάρχει Excel/pywin32: {e}"}

    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        xl.IgnoreRemoteRequests = True
    except Exception:
        pass

    report = {"status": "FAIL", "errors": [], "lines": [], "checks": [], "b2_notes": {}, "s1_notes": {}}
    wb = None
    try:
        wb = xl.Workbooks.Open(os.path.abspath(path))
        wb.Application.CalculateFull()
        for sh in (S1, S2, S3):
            ws = wb.Worksheets(sh)
            ur = ws.UsedRange
            vals = ur.Value
            r0, c0 = ur.Row, ur.Column
            if not isinstance(vals, tuple):
                vals = ((vals,),)
            for ri, row in enumerate(vals):
                for ci, v in enumerate(row):
                    if isinstance(v, int) and v in XL_ERRORS:
                        report["errors"].append(f"{sh}!{get_column_letter(c0 + ci)}{r0 + ri} = {XL_ERRORS[v]}")

        def val(name):
            loc = layout[name]
            return wb.Worksheets(loc["sheet"]).Range(loc["cell"]).Value

        for key in line_keys:
            desc_cell = layout[f"line.{key}.qty"]["cell"]
            row = int(re.sub(r"[A-Z]", "", desc_cell))
            ws1 = wb.Worksheets(S1)
            item = {"key": key, "desc": ws1.Range(f"B{row}").Value, "qty": val(f"line.{key}.qty"),
                    "unit": ws1.Range(f"D{row}").Value, "price": val(f"line.{key}.price"),
                    "total": val(f"line.{key}.total")}
            report["lines"].append(item)
            exp = ref["lines"][key]
            ok = isinstance(item["total"], float) and abs(item["total"] - exp) < 0.01
            report["checks"].append({"check": f"line {key}", "excel": item["total"], "python": exp, "ok": ok})
        for n_key in layout:
            if n_key.startswith("b2.") and n_key.endswith(".note"):
                report["b2_notes"][n_key] = val(n_key)
            elif n_key.startswith("s1.note."):
                report["s1_notes"][n_key] = val(n_key)
        gt = val("grand_total")
        b2_gt = val("b2.grand_total")
        s3_gt = val("s3.grand_total")
        report["grand_total"] = gt
        report["b2_grand_total"] = b2_gt
        report["s3_grand_total"] = s3_gt
        report["n_piles"] = val("n_piles")
        ok = isinstance(gt, float) and abs(gt - ref["grand_total"]) < 0.01
        ok_b2 = isinstance(b2_gt, float) and abs(b2_gt - ref["grand_total"]) < 0.01
        ok_s3 = isinstance(s3_gt, float) and abs(s3_gt - ref["grand_total"]) < 0.01
        report["checks"].append({"check": "grand_total", "excel": gt, "python": ref["grand_total"], "ok": ok})
        report["checks"].append({"check": "b2.grand_total", "excel": b2_gt, "python": ref["grand_total"], "ok": ok_b2})
        report["checks"].append({"check": "s3.grand_total", "excel": s3_gt, "python": ref["grand_total"], "ok": ok_s3})
        if not report["errors"] and all(c["ok"] for c in report["checks"]):
            report["status"] = "PASS"
            wb.Save()
        wb.Close(False)
    finally:
        try:
            xl.IgnoreRemoteRequests = False
            xl.Quit()
        except Exception:
            pass
    return report


# =============================================================================
# 8. CLI
# =============================================================================
def safe_path(p):
    if not os.path.exists(p):
        return p
    base, ext = os.path.splitext(p)
    base = re.sub(r"_v\d+$", "", base)
    k = 2
    while os.path.exists(f"{base}_v{k}{ext}"):
        k += 1
    return f"{base}_v{k}{ext}"


def main(argv=None):
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Μηχανή κοστολόγησης πασσάλων v5.7")
    ap.add_argument("spec")
    ap.add_argument("--out")
    ap.add_argument("--no-verify", action="store_true")
    a = ap.parse_args(argv)
    with open(a.spec, "r", encoding="utf-8") as f:
        spec = json.load(f)
    name = spec.get("project", {}).get("folder") or spec.get("project", {}).get("name", "PROJECT").replace(" ", "_")
    out = a.out or os.path.join(PROJECTS_ROOT, name, f"{name}_Κοστολόγηση_Πασσάλων.xlsx")
    out = safe_path(out)
    try:
        layout, ref, line_keys, sources = build(spec, out)
    except MissingMaterialError as e:
        print("MISSING_MATERIAL")
        print(str(e))
        return 3
    except SpecError as e:
        print(f"SPEC_ERROR: {e}")
        return 5
    print(f"Δημιουργήθηκε: {out}")
    if a.no_verify:
        print("UNVERIFIED (--no-verify)")
        return 2
    rep = verify(out, layout, ref, line_keys)
    print(json.dumps({"file": out, **rep, "python_reference": {
        k: v for k, v in ref.items() if k not in ("lines", "group_calcs")}}, ensure_ascii=False, indent=2, default=str))
    return {"PASS": 0, "UNVERIFIED": 2}.get(rep["status"], 4)


if __name__ == "__main__":
    sys.exit(main())
