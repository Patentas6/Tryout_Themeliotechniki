import math
import sys
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def generate_bolton_tower_excel(target_path):
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # STYLES DEFINITION
    # -------------------------------------------------------------
    font_main_title = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
    fill_navy = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Navy
    
    font_sub_title = Font(name='Calibri', size=11, bold=True, color='1E40AF')
    font_meta = Font(name='Calibri', size=9, italic=True, color='4B5563')
    
    # Red Alert Banner
    fill_red_alert = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    font_red_alert = Font(name='Calibri', size=10, bold=True, color='991B1B')
    border_red_alert = Border(
        left=Side(style='medium', color='DC2626'),
        right=Side(style='medium', color='DC2626'),
        top=Side(style='medium', color='DC2626'),
        bottom=Side(style='medium', color='DC2626')
    )
    
    # Emerald Table (Section 2)
    fill_emerald = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')
    font_emerald = Font(name='Calibri', size=10, bold=True, color='065F46')
    font_emerald_title = Font(name='Calibri', size=11, bold=True, color='065F46')
    border_emerald = Border(
        left=Side(style='thin', color='10B981'),
        right=Side(style='thin', color='10B981'),
        top=Side(style='thin', color='10B981'),
        bottom=Side(style='thin', color='10B981')
    )
    
    # Table Header
    fill_header = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
    font_header = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    
    # Yellow Unit Prices
    fill_yellow = PatternFill(start_color='FEF9C3', end_color='FEF9C3', fill_type='solid')
    font_yellow = Font(name='Calibri', size=10, bold=True, color='854D0E')
    
    # Grand Total
    fill_gold = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    font_gold = Font(name='Calibri', size=11, bold=True, color='92400E')
    border_double_bottom = Border(
        left=Side(style='thin', color='D97706'),
        right=Side(style='thin', color='D97706'),
        top=Side(style='thin', color='D97706'),
        bottom=Side(style='double', color='92400E')
    )
    
    border_box = Border(
        left=Side(style='thin', color='9CA3AF'),
        right=Side(style='thin', color='9CA3AF'),
        top=Side(style='thin', color='9CA3AF'),
        bottom=Side(style='thin', color='9CA3AF')
    )
    
    font_bold = Font(name='Calibri', size=10, bold=True, color='111827')
    font_regular = Font(name='Calibri', size=10, color='374151')
    font_note = Font(name='Calibri', size=9, italic=True, color='6B7280')
    
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    align_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # -------------------------------------------------------------
    # 1. SETUP SHEET 2 FIRST: Τεχνικοί Υπολογισμοί
    # -------------------------------------------------------------
    ws2 = wb.active
    ws2.title = 'Τεχνικοί Υπολογισμοί'
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells('A1:F1')
    ws2['A1'] = 'ΤΕΧΝΙΚΟΙ ΥΠΟΛΟΓΙΣΜΟΙ & ΤΕΚΜΗΡΙΩΣΗ ΠΟΣΟΤΗΤΩΝ — BOLTON TOWER'
    ws2['A1'].font = font_main_title
    ws2['A1'].fill = fill_navy
    ws2['A1'].alignment = align_center
    ws2.row_dimensions[1].height = 26
    
    ws2_col_widths = {'A': 8, 'B': 44, 'C': 15, 'D': 18, 'E': 22, 'F': 45}
    for col_l, w in ws2_col_widths.items():
        ws2.column_dimensions[col_l].width = w

    # Section 1: Inputs
    ws2.merge_cells('A3:F3')
    ws2['A3'] = '1. ΔΕΔΟΜΕΝΑ ΠΟΥ ΔΙΑΒΑΣΤΗΚΑΝ ΑΠΟ ΤΟ ΣΧΕΔΙΟ & ΥΨΟΜΕΤΡΑ (3 ΚΑΘΕΤΟΙ ΤΟΜΕΙΣ)'
    ws2['A3'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    ws2['A3'].fill = fill_header
    ws2['A3'].alignment = align_left
    ws2.row_dimensions[3].height = 22
    
    ws2.cell(4, 1, 'Α/Α').font = font_bold
    ws2.cell(4, 2, 'Παράμετρος Μελέτης').font = font_bold
    ws2.cell(4, 3, 'Μονάδα').font = font_bold
    ws2.cell(4, 4, 'Τιμή (Αριθμός)').font = font_bold
    ws2.cell(4, 5, 'Έλεγχος').font = font_bold
    ws2.cell(4, 6, 'Παρατηρήσεις / Πηγή Σχεδίου').font = font_bold
    for c in range(1, 7):
        ws2.cell(4, c).alignment = align_center if c in (1, 3, 4, 5) else align_left
        ws2.cell(4, c).border = border_box
        ws2.cell(4, c).fill = PatternFill(start_color='E5E7EB', end_color='E5E7EB', fill_type='solid')

    inputs_s2 = [
        (5, '1.1', 'ΤΟΜΕΑΣ 1: ΜΕΤΑΚΙΝΗΣΕΙΣ ΓΕΩΤΡΥΠΑΝΟΥ', '', '', '', 'Έδρα: Άγιος Συλάς, Λεμεσός'),
        (6, '1.1.1', 'Αφετηρία & Τοποθεσία Έργου', 'km', 25.0, 'Ζώνη 1 (0-50km)', 'Εντός πόλης Λεμεσού'),
        (7, '1.1.2', 'Συνολικό Πλήθος Πασσάλων Έργου', 'τεμ.', 50, 'P1-P50', '50 πάσσαλοι βάσει κάτοψης ΘΕΜ/1'),
        (8, '1.1.3', 'Αριθμός Μετακινήσεων Γεωτρύπανου', 'τεμ.', 1, 'Εφάπαξ', '1 πλήρης εγκατάσταση στο εργοτάξιο'),
        (9, '1.1.4', 'Συνολικό Εφάπαξ Ποσό Μετακίνησης', '€', 500.00, '500 €', 'Βασική τιμή καταλόγου (Ζώνη 1)'),
        (10, '1.1.5', 'Επιμερισμένη Τιμή Μετακίνησης ανά Πάσσαλο', '€/πάσσαλο', '=D9/D7', '10.00 €/πάσσαλο', '=D9/D7 (500 € / 50 πάσσαλοι)'),
        
        (11, '1.2', 'ΤΟΜΕΑΣ 2: ΓΕΩΜΕΤΡΙΑ, ΤΡΥΠΗΜΑ & ΥΨΟΜΕΤΡΑ', '', '', '', 'Σχέδιο ΘΕΜ/1'),
        (12, '1.2.1', 'Διάμετρος Πασσάλων Πύργου (D1)', 'mm', 1500, 'Ø1500mm', '21 πάσσαλοι (P1-P21)'),
        (13, '1.2.2', 'Διάμετρος Περιμετρικών Πασσάλων (D2)', 'mm', 1200, 'Ø1200mm', '29 πάσσαλοι (P22-P50)'),
        (14, '1.2.3', 'Στάθμη Πλατφόρμας Γεωτρύπανου (WPL)', 'm', 0.00, 'Εργοτάξιο', 'Δεν αναγράφεται στο σχέδιο'),
        (15, '1.2.4', 'Στάθμη Αποκοπής Πασσάλων (Cut-Off)', 'm', 0.00, 'Πέδιλα B3', 'Δεν αναγράφεται στο σχέδιο'),
        (16, '1.2.5', 'Διαφορά Υψομέτρου Τρυπήματος (Δh = WPL - CutOff)', 'm', '=D14-D15', 'RED ALERT', 'ΚΟΚΚΙΝΟΣ ΣΥΝΑΓΕΡΜΟΣ: Προσωρινά 0.00m'),
        (17, '1.2.6', 'Ύψος Υπερκάλυψης Σκυροδέματος (h_overc)', 'm', '=IF(D16>0,MIN(2.0,D16),0.0)', 'Overconcreting', 'Έως +2.0m αν Δh > 0'),
        (18, '1.2.7', 'Κατηγορία Σκυροδέματος Μελέτης', 'class', 'C30/37', '115,00 €/m³', 'Default πρότυπο εταιρείας'),
        
        (19, '1.3', 'ΤΟΜΕΑΣ 3: ΣΤΑΘΕΡΕΣ ΚΛΩΒΟΥ ΟΠΛΙΣΜΟΥ', '', '', '', 'Τεχνικές προδιαγραφές'),
        (20, '1.3.1', 'Επικάλυψη Οπλισμού (Concrete Cover)', 'mm', 75, '75mm / πλευρά', 'Ελάχιστη επικάλυψη για bored piles'),
        (21, '1.3.2', 'Αναμονή Starter στο πέδιλο (L_starter)', 'm', 0.50, '0.50m', 'Αναμονή σύνδεσης με θεμελίωση B3'),
        (22, '1.3.3', 'Γωνία Άνω Σύνδεσης (L_top_bend)', 'm', 0.50, '0.50m', 'Τυπική γωνία αγκύρωσης κορυφής'),
        (23, '1.3.4', 'Γωνία Πυθμένα (L_bottom_bend)', 'm', 0.00, '0.00m', 'Δεν δίνεται στο σχέδιο / Δεν απαιτείται'),
        (24, '1.3.5', 'Διάμετρος Σούστας / Σπείρας (Y_spiral)', 'mm', 10, 'Y10', 'Σπείρα εγκάρσιου οπλισμού'),
        (25, '1.3.6', 'Βήμα Σούστας (Pitch s)', 'mm', 150, 's = 150mm', 'Τυπικό βήμα σούστας')
    ]
    
    for item in inputs_s2:
        r, aa, desc, unit, val, check, comm = item
        ws2.row_dimensions[r].height = 20
        is_sector_header = desc.startswith('ΤΟΜΕΑΣ')
        
        c1 = ws2.cell(r, 1, aa)
        c2 = ws2.cell(r, 2, desc)
        c3 = ws2.cell(r, 3, unit)
        c4 = ws2.cell(r, 4, val)
        c5 = ws2.cell(r, 5, check)
        c6 = ws2.cell(r, 6)
        c6.value = comm
        c6.data_type = 's'
        
        c1.alignment = align_center
        c2.alignment = align_left
        c3.alignment = align_center
        c4.alignment = align_right if isinstance(val, (int, float)) or (isinstance(val, str) and val.startswith('=')) else align_center
        c5.alignment = align_center
        c6.alignment = align_left
        
        if is_sector_header:
            for c in range(1, 7):
                cell = ws2.cell(r, c)
                cell.font = font_bold
                cell.fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
                cell.border = border_box
        else:
            for c in range(1, 7):
                ws2.cell(r, c).font = font_regular if c != 4 else font_bold
                ws2.cell(r, c).border = border_box
                
            if isinstance(val, (int, float)) or (isinstance(val, str) and val.startswith('=')):
                if unit in ('m', 'm³'):
                    c4.number_format = '#,##0.00'
                elif unit in ('mm', 'τεμ.'):
                    c4.number_format = '#,##0'
                elif unit in ('€', '€/πάσσαλο'):
                    c4.number_format = '#,##0.00 "€"'
                    
            if 'Δh' in desc:
                c4.fill = fill_red_alert
                c4.font = font_red_alert
                c4.border = border_red_alert
                c5.fill = fill_red_alert
                c5.font = font_red_alert
                c5.border = border_red_alert

    # Section 3: Constants (Rows 36 to 42)
    # We place Section 3 first in code so we know its rows (36 to 42)
    ws2.merge_cells('A36:F36')
    ws2['A36'] = '3. ΣΤΑΘΕΡΕΣ ΤΕΧΝΙΚΕΣ ΠΑΡΑΔΟΧΕΣ & ΦΥΣΙΚΕΣ ΣΤΑΘΕΡΕΣ (ΣΤΗΛΗ D = ΑΡΙΘΜΟΙ ΜΟΝΟ)'
    ws2['A36'].font = Font(name='Calibri', size=11, bold=True, color='1E40AF')
    ws2['A36'].fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
    ws2['A36'].alignment = align_left
    ws2.row_dimensions[36].height = 22
    
    constants_data = [
        (37, '3.1', 'Συντελεστής Απώλειας / Υπερκατανάλωσης Σκυροδέματος', '%', 0.10, '10%', 'Πάγια προσαύξηση όγκου +10% (V_conc = V_exc × 1.10)'),
        (38, '3.2', 'Ειδικό Βάρος Δομικού Χάλυβα Οπλισμού (ρ)', 'kg/m³', 7850, '7850 kg/m³', 'Φυσική σταθερά χάλυβα B500C'),
        (39, '3.3', 'Μέγιστο Εμπορικό Μήκος Βέργας Χάλυβα', 'm', 12.00, '12.00 m', 'Όριο πέραν του οποίου υπολογίζεται μάτιση'),
        (40, '3.4', 'Συντελεστής Μήκους Μάτισης Ράβδων (Lapping Factor)', 'x Φ', 50, '50 × Φ', 'Κανόνας 50 φορές τη διάμετρο της ράβδου'),
        (41, '3.5', 'Μέγιστο Ύψος Υπερκάλυψης Σκυροδέματος (+2.0m)', 'm', 2.00, '2.00 m', 'Μέγιστο όριο υπερύψωσης σκυροδέματος πάνω από cut-off'),
        (42, '3.6', 'Μαθηματική Σταθερά π', 'rad', '=PI()', '3.14159...', 'Σταθερά κύκλου Excel =PI()')
    ]
    
    for r, aa, desc, unit, val, check, comm in constants_data:
        ws2.row_dimensions[r].height = 20
        c1 = ws2.cell(r, 1, aa)
        c2 = ws2.cell(r, 2, desc)
        c3 = ws2.cell(r, 3, unit)
        c4 = ws2.cell(r, 4, val)
        c5 = ws2.cell(r, 5, check)
        c6 = ws2.cell(r, 6)
        c6.value = comm
        c6.data_type = 's'
        
        c1.alignment = align_center
        c2.alignment = align_left
        c3.alignment = align_center
        c4.alignment = align_right
        c5.alignment = align_center
        c6.alignment = align_left
        
        if unit == '%':
            c4.number_format = '0.0%'
        elif unit == 'kg/m³':
            c4.number_format = '#,##0'
        elif unit == 'm':
            c4.number_format = '0.00'
        elif unit == 'x Φ':
            c4.number_format = '0'
        else:
            c4.number_format = '0.00000'
            
        for c in range(1, 7):
            cell = ws2.cell(r, c)
            cell.font = font_regular if c != 4 else font_bold
            cell.border = border_box

    # Section 4: Analytical Calculations for 6 Pile Types
    # Starts at row 44
    ws2.merge_cells('A44:F44')
    ws2['A44'] = '4. ΑΝΑΛΥΤΙΚΑ ΜΑΘΗΜΑΤΙΚΑ ΒΗΜΑ-ΒΗΜΑ & ΔΙΑΧΩΡΙΣΜΟΣ 6 ΤΥΠΩΝ ΠΑΣΣΑΛΩΝ'
    ws2['A44'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    ws2['A44'].fill = fill_navy
    ws2['A44'].alignment = align_left
    ws2.row_dimensions[44].height = 24
    
    current_math_r = 45

    pile_types_spec = [
        {
            'title': '4.1 ΤΥΠΟΣ 1: ΠΑΣΣΑΛΟΙ Ø1500mm — ΒΑΘΟΣ 22.00m (21 ΤΕΜΑΧΙΑ: P1 - P21)',
            'count': 21,
            'dia_mm': 1500,
            'l_pile': 22.0,
            'n_vert': 22,
            'y_vert': 28,
            'starter': 0.50,
            'top_bend': 0.50,
            'bot_bend': 0.00,
            'cover_mm': 75,
            'y_spiral': 10,
            'pitch_mm': 150
        },
        {
            'title': '4.2 ΤΥΠΟΣ 2: ΠΑΣΣΑΛΟΙ Ø1200mm — ΒΑΘΟΣ 19.00m (6 ΤΕΜΑΧΙΑ: P43 - P48)',
            'count': 6,
            'dia_mm': 1200,
            'l_pile': 19.0,
            'n_vert': 18,
            'y_vert': 25,
            'starter': 0.50,
            'top_bend': 0.50,
            'bot_bend': 0.00,
            'cover_mm': 75,
            'y_spiral': 10,
            'pitch_mm': 150
        },
        {
            'title': '4.3 ΤΥΠΟΣ 3: ΠΑΣΣΑΛΟΙ Ø1200mm — ΒΑΘΟΣ 18.00m (3 ΤΕΜΑΧΙΑ: P33 - P35)',
            'count': 3,
            'dia_mm': 1200,
            'l_pile': 18.0,
            'n_vert': 16,
            'y_vert': 25,
            'starter': 0.50,
            'top_bend': 0.50,
            'bot_bend': 0.00,
            'cover_mm': 75,
            'y_spiral': 10,
            'pitch_mm': 150
        },
        {
            'title': '4.4 ΤΥΠΟΣ 4: ΠΑΣΣΑΛΟΙ Ø1200mm — ΒΑΘΟΣ 17.00m (5 ΤΕΜΑΧΙΑ: P36 - P40)',
            'count': 5,
            'dia_mm': 1200,
            'l_pile': 17.0,
            'n_vert': 16,
            'y_vert': 25,
            'starter': 0.50,
            'top_bend': 0.50,
            'bot_bend': 0.00,
            'cover_mm': 75,
            'y_spiral': 10,
            'pitch_mm': 150
        },
        {
            'title': '4.5 ΤΥΠΟΣ 5: ΠΑΣΣΑΛΟΙ Ø1200mm — ΒΑΘΟΣ 16.00m (13 ΤΕΜΑΧΙΑ: P22-P32, P41-P42)',
            'count': 13,
            'dia_mm': 1200,
            'l_pile': 16.0,
            'n_vert': 14,
            'y_vert': 25,
            'starter': 0.50,
            'top_bend': 0.50,
            'bot_bend': 0.00,
            'cover_mm': 75,
            'y_spiral': 10,
            'pitch_mm': 150
        },
        {
            'title': '4.6 ΤΥΠΟΣ 6: ΠΑΣΣΑΛΟΙ Ø1200mm — ΒΑΘΟΣ 10.50m (2 ΤΕΜΑΧΙΑ: P49 - P50)',
            'count': 2,
            'dia_mm': 1200,
            'l_pile': 10.5,
            'n_vert': 14,
            'y_vert': 20,
            'starter': 0.50,
            'top_bend': 0.50,
            'bot_bend': 0.00,
            'cover_mm': 75,
            'y_spiral': 10,
            'pitch_mm': 150
        }
    ]

    # Store row locations of summary outputs for Section 2 links
    # For each type: {drill_row, conc_row, steel_row}
    type_summary_cells = []

    for type_idx, pt in enumerate(pile_types_spec, 1):
        ws2.merge_cells(f'A{current_math_r}:F{current_math_r}')
        ws2[f'A{current_math_r}'] = pt['title']
        ws2[f'A{current_math_r}'].font = font_sub_title
        ws2[f'A{current_math_r}'].fill = PatternFill(start_color='E0E7FF', end_color='E0E7FF', fill_type='solid')
        ws2[f'A{current_math_r}'].alignment = align_left
        ws2.row_dimensions[current_math_r].height = 22
        current_math_r += 1
        
        ws2.cell(current_math_r, 1, 'Βήμα').font = font_bold
        ws2.cell(current_math_r, 2, 'Μαθηματικό Βήμα / Υπολογιζόμενο Μέγεθος').font = font_bold
        ws2.cell(current_math_r, 3, 'Μονάδα').font = font_bold
        ws2.cell(current_math_r, 4, 'Υπολογισμός (Φόρμουλα Excel)').font = font_bold
        ws2.cell(current_math_r, 5, 'Σύνολο Ομάδας').font = font_bold
        ws2.cell(current_math_r, 6, 'Επεξήγηση Τύπου (Note)').font = font_bold
        for c in range(1, 7):
            ws2.cell(current_math_r, c).alignment = align_center if c in (1, 3, 4, 5) else align_left
            ws2.cell(current_math_r, c).border = border_box
            ws2.cell(current_math_r, c).fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
        current_math_r += 1

        r_start = current_math_r
        
        # Save summary rows
        # Step 4 is at r_start + 3 (Total meters in E)
        # Step 6 is at r_start + 5 (Total concrete in E)
        # Step 17 is at r_start + 16 (Total steel in E)
        type_summary_cells.append({
            'drill_e': f'E{r_start+3}',
            'drill_d': f'D{r_start+3}',
            'conc_e': f'E{r_start+5}',
            'conc_d': f'D{r_start+5}',
            'steel_e': f'E{r_start+16}',
            'steel_d': f'D{r_start+16}',
            'dia': pt['dia_mm'],
            'count': pt['count']
        })
        
        steps = [
            ('1', 'Πλήθος πασσάλων ομάδας', 'τεμ.', pt['count'], f"{pt['count']} πάσσαλοι", 'Διαβάστηκε από την κάτοψη'),
            ('2', 'Ονομαστική διάμετρος πασσάλου (D)', 'mm', pt['dia_mm'], f"Ø{pt['dia_mm']}mm", 'Ονομαστική διάμετρος οπής'),
            ('3', 'Καθαρό βάθος πασσάλου (L_pile)', 'm', pt['l_pile'], f"{pt['l_pile']:.2f} m", 'Βάθος διάτρησης στο έδαφος'),
            ('4', 'Συνολικό βάθος διάτρησης (L_drill = L_pile + Δh)', 'm', f'=D{r_start+2}+D16', f'=D{r_start+3}*D{r_start}', 'L_drill = L_pile + Δh (Σύνολο = L_drill × Πλήθος)'),
            ('5', 'Γεωμετρικός όγκος σκυροδέματος ανά πάσσαλο', 'm³', f'=D42*(D{r_start+1}/1000)^2/4*(D{r_start+2}+D17)', f'=D{r_start+4}*D{r_start}', 'V_geom = π × D²/4 × (L_pile + h_overc)'),
            ('6', 'Όγκος σκυροδέματος με +10% απώλεια (V_conc)', 'm³', f'=D{r_start+4}*(1+D37)', f'=D{r_start+5}*D{r_start}', 'V_conc = V_geom × (1 + Loss 10%)'),
            ('7', 'Κάθετος οπλισμός: Πλήθος & Διατομή', 'ράβδοι', pt['n_vert'], f"{pt['n_vert']}Y{pt['y_vert']}", f"{pt['n_vert']} ράβδοι διατομής Y{pt['y_vert']}"),
            ('8', 'Βασικό μήκος ράβδου (L_base = L + starter + bends)', 'm', f"=D{r_start+2}+{pt['starter']}+{pt['top_bend']}+{pt['bot_bend']}", f"{pt['l_pile'] + pt['starter'] + pt['top_bend']:.2f} m", 'L_base = L_pile + starter + top_bend + bot_bend'),
            ('9', 'Αριθμός ματίσεων ανά ράβδο (Lapping count)', 'τεμ.', f'=IF(D{r_start+7}>D39, INT((D{r_start+7}-0.001)/D39), 0)', 'Ματίσεις', '=IF(L_base > 12m, INT((L_base - 0.001)/12), 0)'),
            ('10', 'Μήκος μάτισης (50Φ)', 'm', f"=D40*({pt['y_vert']}/1000)", f"{50*pt['y_vert']/1000:.2f} m", 'Lap_len = 50 × (Y_vert / 1000)'),
            ('11', 'Ανεπτυγμένο μήκος κάθετης ράβδου (L_bar)', 'm', f'=D{r_start+7}+(D{r_start+8}*D{r_start+9})', f'=D{r_start+10}*{pt["n_vert"]}', 'L_bar = L_base + (Laps × Lap_length)'),
            ('12', f"Μοναδιαίο βάρος κάθετης ράβδου Y{pt['y_vert']}", 'kg/m', f"=D42*({pt['y_vert']}/1000)^2/4*D38", f"{math.pi * ((pt['y_vert']/1000)**2) / 4 * 7850:.3f} kg/m", 'w = π × d²/4 × 7850 kg/m³'),
            ('13', 'Συνολικό βάρος κάθετου οπλισμού ανά πάσσαλο', 'kg', f'=D{r_start+6}*D{r_start+10}*D{r_start+11}', f'=D{r_start+12}*D{r_start}', 'W_vert = N_vert × L_bar × w_vert'),
            ('14', 'Διάμετρος & Περίμετρος σπείρας κλωβού', 'm', f'=D42*(D{r_start+1}-2*{pt["cover_mm"]})/1000', 'Περίμετρος', 'D_cage = D - 2×Cover, Περίμετρος = π × D_cage'),
            ('15', 'Αριθμός σπειρών σούστας ανά πάσσαλο', 'σπείρες', f'=ROUNDUP(D{r_start+2}/({pt["pitch_mm"]}/1000), 0)', 'Σπείρες', 'Turns = ROUNDUP(L_pile / Pitch, 0)'),
            ('16', f"Βάρος σούστας Y{pt['y_spiral']}/{pt['pitch_mm']} ανά πάσσαλο", 'kg', f"=D{r_start+14}*D{r_start+13}*(D42*({pt['y_spiral']}/1000)^2/4*D38)", f'=D{r_start+15}*D{r_start}', 'W_spiral = Turns × Περίμετρος × w_spiral'),
            ('17', 'ΣΥΝΟΛΙΚΟ ΒΑΡΟΣ ΧΑΛΥΒΑ (ΚΑΘΕΤΑ + ΣΟΥΣΤΑ)', 'kg', f'=D{r_start+12}+D{r_start+15}', f'=D{r_start+16}*D{r_start}', 'W_steel_tot = W_vert + W_spiral')
        ]
        
        for st_idx, (b_num, b_desc, b_unit, b_form_pile, b_form_tot, b_note) in enumerate(steps):
            r_cur = r_start + st_idx
            ws2.row_dimensions[r_cur].height = 20
            
            c1 = ws2.cell(r_cur, 1, b_num)
            c2 = ws2.cell(r_cur, 2, b_desc)
            c3 = ws2.cell(r_cur, 3, b_unit)
            c4 = ws2.cell(r_cur, 4, b_form_pile)
            c5 = ws2.cell(r_cur, 5, b_form_tot)
            c6 = ws2.cell(r_cur, 6)
            c6.value = b_note
            c6.data_type = 's'
            
            c1.alignment = align_center
            c2.alignment = align_left
            c3.alignment = align_center
            c4.alignment = align_right
            c5.alignment = align_right
            c6.alignment = align_left
            
            is_summary_step = b_num in ('4', '6', '17')
            for col_idx in range(1, 7):
                cell = ws2.cell(r_cur, col_idx)
                cell.font = font_bold if is_summary_step else font_regular
                cell.border = border_box
                if is_summary_step:
                    cell.fill = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
                    
            if b_unit in ('m', 'm³', 'kg'):
                c4.number_format = '#,##0.00'
                if str(b_form_tot).startswith('='):
                    c5.number_format = '#,##0.00'
            elif b_unit in ('mm', 'τεμ.', 'σπείρες', 'ράβδοι'):
                c4.number_format = '#,##0'
            elif b_unit == 'kg/m':
                c4.number_format = '#,##0.000'
                
        current_math_r = r_start + len(steps) + 1

    # Now that we know all summary cell references, let's write Section 2 (Emerald Table)
    # Emerald Table is at rows 27 to 34
    ws2.merge_cells('A27:F27')
    ws2['A27'] = '2. ΤΕΛΙΚΕΣ ΠΟΣΟΤΗΤΕΣ ΓΙΑ ΤΗΝ ΟΙΚΟΝΟΜΙΚΗ ΠΡΟΣΦΟΡΑ (SUMMARY RESULTS)'
    ws2['A27'].font = font_emerald_title
    ws2['A27'].fill = fill_emerald
    ws2['A27'].alignment = align_left
    ws2.row_dimensions[27].height = 24
    
    sec2_headers = ['Α/Α', 'Περιγραφή Τελικής Ποσότητας Έργου', 'Μονάδα', 'Ποσότητα ανά Πάσσαλο', 'Συνολική Ποσότητα', 'Σύνδεση με Προσφορά & Τεκμηρίωση']
    ws2.row_dimensions[28].height = 22
    for c_idx, h_text in enumerate(sec2_headers, 1):
        cell = ws2.cell(28, c_idx, h_text)
        cell.font = font_emerald
        cell.fill = fill_emerald
        cell.border = border_emerald
        cell.alignment = align_center if c_idx in (1, 3, 4, 5) else align_left

    # Construct the exact sums:
    # Drill Ø1500: type 1 only (type_summary_cells[0]['drill_e'])
    # Drill Ø1200: types 2 to 6 (sum of drill_e for types 1..5 in 0-indexed list)
    dia1200_drill_cells = ",".join(t['drill_e'] for t in type_summary_cells if t['dia'] == 1200)
    all_conc_cells = ",".join(t['conc_e'] for t in type_summary_cells)
    all_steel_cells = ",".join(t['steel_e'] for t in type_summary_cells)

    emerald_rows = [
        (29, '2.1', 'Τρέχοντα Μέτρα Διάτρησης Ø1500mm (21 πάσσαλοι)', 'm',
         '=E29/21', f"={type_summary_cells[0]['drill_e']}", 'Τροφοδοτεί Γραμμή 2a Φύλλου 1 (21 τεμ. × 22.00m)'),
         
        (30, '2.2', 'Τρέχοντα Μέτρα Διάτρησης Ø1200mm (29 πάσσαλοι)', 'm',
         '=E30/29', f"=SUM({dia1200_drill_cells})", 'Τροφοδοτεί Γραμμή 2b Φύλλου 1 (29 τεμ., 10.5-19.0m)'),
         
        (31, '2.3', 'Κυβικά Σκυροδέματος C30/37 με +10% απώλεια', 'm³',
         '=E31/D7', f"=SUM({all_conc_cells})", 'Τροφοδοτεί Γραμμή 3 Φύλλου 1 (V_conc με 10% loss)'),
         
        (32, '2.4', 'Συνολικό Βάρος Χάλυβα Οπλισμού B500C', 'kg',
         '=E32/D7', f"=SUM({all_steel_cells})", 'Τροφοδοτεί Γραμμή 4 Φύλλου 1 (Κάθετες & Σούστα)'),
         
        (33, '2.5', 'Sleeves Προστασίας Starter Αναμονών', 'τεμ.',
         '=IF(D16>0, 18, 0)', '=D33*D7', 'Ενεργοποιούνται αυτόματα εάν Δh > 0'),
         
        (34, '2.6', 'Κεφαλοκοπή Πασσάλων / Pile Breaking', 'τεμ.',
         '=IF(D16>0, 1, 0)', '=D34*D7', 'Ενεργοποιείται αυτόματα εάν Δh > 0')
    ]
    
    for r, aa, desc, unit, form_pile, form_tot, note_txt in emerald_rows:
        ws2.row_dimensions[r].height = 20
        c1 = ws2.cell(r, 1, aa)
        c2 = ws2.cell(r, 2, desc)
        c3 = ws2.cell(r, 3, unit)
        c4 = ws2.cell(r, 4, form_pile)
        c5 = ws2.cell(r, 5, form_tot)
        c6 = ws2.cell(r, 6)
        c6.value = note_txt
        c6.data_type = 's'
        
        c1.alignment = align_center
        c2.alignment = align_left
        c3.alignment = align_center
        c4.alignment = align_right
        c5.alignment = align_right
        c6.alignment = align_left
        
        c4.number_format = '#,##0.00'
        c5.number_format = '#,##0.00'
        
        for col_idx in range(1, 7):
            cell = ws2.cell(r, col_idx)
            cell.font = font_emerald
            cell.fill = fill_emerald
            cell.border = border_emerald


    # -------------------------------------------------------------
    # 2. SETUP SHEET 1: Οικονομική Προσφορά
    # -------------------------------------------------------------
    ws1 = wb.create_sheet('Οικονομική Προσφορά', 0) # Place as first sheet!
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.merge_cells('A1:G1')
    ws1['A1'] = 'ΘΕΜΕΛΙΟΤΕΧΝΙΚΗ ΛΤΔ (Themeliotechniki LTD) — ΕΙΔΙΚΑ ΓΕΩΤΕΧΝΙΚΑ ΕΡΓΑ & ΠΑΣΣΑΛΟΙ'
    ws1['A1'].font = font_main_title
    ws1['A1'].fill = fill_navy
    ws1['A1'].alignment = align_center
    ws1.row_dimensions[1].height = 28
    
    ws1.merge_cells('A2:G2')
    ws1['A2'] = 'ΕΠΙΣΗΜΗ ΟΙΚΟΝΟΜΙΚΗ ΠΡΟΣΦΟΡΑ ΠΑΣΣΑΛΩΝ ΘΕΜΕΛΙΩΣΗΣ — BOLTON TOWER, ΛΕΜΕΣΟΣ'
    ws1['A2'].font = font_sub_title
    ws1['A2'].alignment = align_center
    ws1.row_dimensions[2].height = 20
    
    ws1.merge_cells('A3:G3')
    ws1['A3'] = 'Ημερομηνία: 15/10/2026 | Μελέτη: ΘΕΜ/1 Hyperstatic (220138) | Σύνολο Πασσάλων: 50 τεμ. (P1-P50) | Ζώνη 1: Λεμεσός'
    ws1['A3'].font = font_meta
    ws1['A3'].alignment = align_center
    ws1.row_dimensions[3].height = 18
    
    # RED ALERT BANNER (Rows 4-5)
    ws1.merge_cells('A4:G5')
    banner_text = (
        "🛑 ΠΡΟΣΟΧΗ — ΚΟΚΚΙΝΟΣ ΣΥΝΑΓΕΡΜΟΣ (RED ALERT): Στα δοθέντα σχέδια δεν αναγράφεται υψομετρική διαφορά στάθμης "
        "εργασίας γεωτρύπανου (WPL) έναντι στάθμης αποκοπής (Cut-off Level). Η παρούσα προσφορά υπολογίζει προσωρινά "
        "Δh = 0.00 m (μηδενική κενή διάτρηση, χωρίς sleeves και κεφαλοκοπή). Απαιτείται επιβεβαίωση υψομέτρων στο εργοτάξιο πριν την οριστικοποίηση!"
    )
    ws1['A4'] = banner_text
    ws1['A4'].font = font_red_alert
    ws1['A4'].fill = fill_red_alert
    ws1['A4'].alignment = align_wrap
    
    for r in range(4, 6):
        ws1.row_dimensions[r].height = 20
        for c in range(1, 8):
            cell = ws1.cell(r, c)
            cell.fill = fill_red_alert
            t_side = Side(style='medium', color='DC2626') if r == 4 else None
            b_side = Side(style='medium', color='DC2626') if r == 5 else None
            l_side = Side(style='medium', color='DC2626') if c == 1 else None
            r_side = Side(style='medium', color='DC2626') if c == 7 else None
            cell.border = Border(top=t_side, bottom=b_side, left=l_side, right=r_side)

    # Table Header Row 7
    headers_s1 = [
        ('Α/Α', 6, align_center),
        ('Περιγραφή Εργασίας / Υλικού', 48, align_left),
        ('Ποσότητα ανά Πάσσαλο', 20, align_right),
        ('Μονάδα', 12, align_center),
        ('Τιμή Μονάδας (€)', 18, align_right),
        ('Κόστος ανά Πάσσαλο (€)', 22, align_right),
        ('Σύνολο Όλων των Πασσάλων (€)', 26, align_right)
    ]
    
    ws1.row_dimensions[7].height = 26
    for col_idx, (h_text, width, h_align) in enumerate(headers_s1, 1):
        cell = ws1.cell(7, col_idx, h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_box
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = width

    # Pricing Rows: linking to exact Section 2 Emerald rows (29 to 34)
    pricing_data = [
        (8, '1', 'Μετακινήσεις & Εγκατάσταση Γεωτρύπανου (Lump Sum επιμερισμένο ανά πάσσαλο)',
         1.0, 'πάσσαλος', "='Τεχνικοί Υπολογισμοί'!D10", '=C8*E8', "='Τεχνικοί Υπολογισμοί'!D9"),
        
        (9, '2a', 'Τρύπημα / Διάτρηση Πασσάλων Ø1500mm (21 τεμάχια, Βάθος 22.00m)',
         "='Τεχνικοί Υπολογισμοί'!D29", 'm', 190.00, '=C9*E9', "='Τεχνικοί Υπολογισμοί'!E29*E9"),
         
        (10, '2b', 'Τρύπημα / Διάτρηση Πασσάλων Ø1200mm (29 τεμάχια, Βάθη 10.50m - 19.00m)',
         "='Τεχνικοί Υπολογισμοί'!D30", 'm', 150.00, '=C10*E10', "='Τεχνικοί Υπολογισμοί'!E30*E10"),
         
        (11, '3', 'Σκυροδέτηση Πασσάλων C30/37 (Κουγκρί με +10% συντελεστή απώλειας)',
         "='Τεχνικοί Υπολογισμοί'!D31", 'm³', 115.00, '=C11*E11', "='Τεχνικοί Υπολογισμοί'!E31*E11"),
         
        (12, '4', 'Χάλυβας Οπλισμού B500C (Κάθετες Ράβδοι & Σούστα Y10/150)',
         "='Τεχνικοί Υπολογισμοί'!D32", 'kg', 1.10, '=C12*E12', "='Τεχνικοί Υπολογισμοί'!E32*E12"),
         
        (13, '5', 'Sleeves Προστασίας Starter Αναμονών (Debonding Sleeves)',
         "='Τεχνικοί Υπολογισμοί'!D33", 'τεμ.', 3.50, '=C13*E13', "='Τεχνικοί Υπολογισμοί'!E33*E13"),
         
        (14, '6', 'Κεφαλοκοπή Πασσάλων / Pile Breaking (Σπάσιμο υπερυψωμένου σκυροδέματος)',
         "='Τεχνικοί Υπολογισμοί'!D34", 'πάσσαλος', 60.00, '=C14*E14', "='Τεχνικοί Υπολογισμοί'!E34*E14")
    ]
    
    for r, aa, desc, qty, unit, rate, cost_pile, cost_tot in pricing_data:
        ws1.row_dimensions[r].height = 22
        ws1.cell(r, 1, aa).alignment = align_center
        ws1.cell(r, 2, desc).alignment = align_left
        
        c3 = ws1.cell(r, 3, qty)
        c3.alignment = align_right
        c3.number_format = '#,##0.00' if isinstance(qty, str) or qty != 1.0 else '0'
        
        ws1.cell(r, 4, unit).alignment = align_center
        
        c5 = ws1.cell(r, 5, rate)
        c5.alignment = align_right
        c5.number_format = '#,##0.00 "€"'
        c5.fill = fill_yellow
        c5.font = font_yellow
        
        c6 = ws1.cell(r, 6, cost_pile)
        c6.alignment = align_right
        c6.number_format = '#,##0.00 "€"'
        c6.font = font_bold
        
        c7 = ws1.cell(r, 7, cost_tot)
        c7.alignment = align_right
        c7.number_format = '#,##0.00 "€"'
        c7.font = font_bold
        
        for c in range(1, 8):
            if c != 5:
                ws1.cell(r, c).font = font_regular if c not in (1, 6, 7) else font_bold
            ws1.cell(r, c).border = border_box

    # Grand Total Row 15
    ws1.row_dimensions[15].height = 26
    ws1.merge_cells('A15:E15')
    ws1['A15'] = 'ΓΕΝΙΚΟ ΣΥΝΟΛΟ ΟΙΚΟΝΟΜΙΚΗΣ ΠΡΟΣΦΟΡΑΣ (50 ΠΑΣΣΑΛΟΙ — ΧΩΡΙΣ ΦΠΑ)'
    ws1['A15'].font = font_gold
    ws1['A15'].alignment = Alignment(horizontal='right', vertical='center')
    
    c_tot_pile = ws1.cell(15, 6, "=G15/'Τεχνικοί Υπολογισμοί'!D7")
    c_tot_pile.font = font_gold
    c_tot_pile.alignment = align_right
    c_tot_pile.number_format = '#,##0.00 "€"'
    
    c_tot_all = ws1.cell(15, 7, '=SUM(G8:G14)')
    c_tot_all.font = font_gold
    c_tot_all.alignment = align_right
    c_tot_all.number_format = '#,##0.00 "€"'
    
    for c in range(1, 8):
        cell = ws1.cell(15, c)
        cell.fill = fill_gold
        cell.border = border_double_bottom

    # Notes & Conditions (Rows 17-27)
    ws1.cell(17, 1, 'ΟΡΟΙ & ΠΑΡΑΤΗΡΗΣΕΙΣ ΟΙΚΟΝΟΜΙΚΗΣ ΠΡΟΣΦΟΡΑΣ:').font = font_bold
    notes = [
        "1. Στις παραπάνω τιμές μονάδας και στο συνολικό ποσό ΔΕΝ συμπεριλαμβάνεται Φ.Π.Α. 19%.",
        "2. Τιμολόγηση Διάτρησης κατόπιν οδηγίας πελάτη: Ø1500mm στα 190,00 €/m και Ø1200mm στα 150,00 €/m.",
        "3. Σκυρόδεμα κατηγορίας C30/37 στα 115,00 €/m³ με συνυπολογισμένο πάγιο συντελεστή απώλειας / υπερκατανάλωσης +10%.",
        "4. Χάλυβας οπλισμού ποιότητας B500C ενιαία στα 1,10 €/kg (αγορά, κοπή, κατασκευή κλωβών και τοποθέτηση).",
        "5. Διαχωρισμός Κλωβών ανά τύπο βάθους: Τύπος 1 (22m) με 22Y28, Τύποι 2-5 (16-19m) με 14-18Y25, Τύπος 6 (10.5m) με 14Y20 χωρίς μάτιση.",
        "6. ΚΟΚΚΙΝΟΣ ΣΥΝΑΓΕΡΜΟΣ ΥΨΟΜΕΤΡΩΝ: Επειδή στο σχέδιο δεν δίνεται διαφορά στάθμης διάτρησης και cut-off, έχει τεθεί προσωρινά Δh = 0.00m.",
        "   Εάν κατά την έναρξη του έργου διαπιστωθεί κενή διάτρηση (Δh > 0), θα τιμολογηθούν τα επιπλέον τρέχοντα μέτρα διάτρησης, τα sleeves (3,50 €/τεμ.) και η κεφαλοκοπή (60,00 €/πάσσαλο).",
        "7. Μετακινήσεις: Εφάπαξ ποσό 500,00 € (Ζώνη 1: 0-50km από έδρα Αγίου Συλά, Λεμεσός), επιμερισμένο ισόποσα σε 10,00 € ανά πάσσαλο.",
        "8. Πληρωμή: Έναντι πιστοποιήσεων προόδου εργασιών με τελική επιμέτρηση πραγματικών μέτρων διάτρησης και σκυροδέτησης."
    ]
    for idx, n_text in enumerate(notes, 18):
        ws1.row_dimensions[idx].height = 18
        cell = ws1.cell(idx, 1, n_text)
        cell.font = font_note
        cell.alignment = align_left

    # Save to file
    wb.save(target_path)
    print(f"Generated successfully: {target_path}")

if __name__ == '__main__':
    target = r"C:\Users\nikth\Desktop\NIKOLAS TRY\Projects\Piles\BOLTON_TOWER\BOLTON_TOWER_Κοστολόγηση_Πασσάλων.xlsx"
    generate_bolton_tower_excel(target)
    target_std = r"C:\Users\nikth\Desktop\NIKOLAS TRY\Projects\Piles\BOLTON_TOWER\Pile_Quotation_2Sheets.xlsx"
    generate_bolton_tower_excel(target_std)
